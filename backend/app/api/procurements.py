"""Procurement API."""
import json
import math
import re
import shutil
import logging
from fastapi import APIRouter, Depends, HTTPException, Query

logger = logging.getLogger("shanhai")
from sqlalchemy.orm import Session
from datetime import date
from app.database import get_db
from app.config import ARCHIVED_FILE_ROOT
from app.models.user import User
from app.models.project import Project
from app.models.procurement import Procurement
from app.models.supplier import Supplier
from app.models.ledger import Ledger
from app.models.process_file_sync_status import ProcessFileSyncStatus
from app.core.auth import get_current_user
from app.schemas.procurement import ProcurementCreate, ProcurementUpdate, ProcurementRemarkPatch
from app.services.project_service import is_officer
from app.services.archive_service import get_archive_contract_folder, find_dual_sibling as resolve_dual_sibling
from app.services.procurement_service import (
    add_procurement_ids_to_folder_marker,
    create_procurement_folder,
    delete_procurement_resources,
    rename_procurement_folders,
    sync_ledgers_on_procurement_update,
)
from app.services.process_file_sync_service import (
    ensure_sync_status_records,
    get_primary_procurement_id,
    reset_sync_status_records_for_new_procurement,
    snap_sync_status_mtimes_from_disk,
    sync_process_files_on_form_save,
)
from app.services.word_service import get_template_path, build_context, render_docx
from app.services.emit_event import emit
from app.events_schema import EventType
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn


def _save_mulu_docx(folder_path, title: str, time_records: list) -> None:
    """生成目录.docx：标题宋体四号居中；下方两列表格（流程节点 | 日期）。"""
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(14)  # 四号
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr_list = list(time_records or [])
    n = len(tr_list)
    table = doc.add_table(rows=max(1, 1 + n), cols=2)
    table.style = "Table Grid"
    h0, h1 = table.rows[0].cells[0], table.rows[0].cells[1]
    h0.text = "流程节点"
    h1.text = "日期"
    for i, tr in enumerate(tr_list, start=1):
        row = table.rows[i].cells
        row[0].text = str(tr.get("flow_name", "") or "")
        row[1].text = str(tr.get("date_val", "") or "")
    doc.save(folder_path / "目录.docx")


router = APIRouter(prefix="/api/procurements", tags=["procurements"])


def check_permission(project: Project, current_user: User):
    if current_user.role != "系统管理员" and not is_officer(project, current_user.id):
        raise HTTPException(status_code=403, detail="无编辑权限")


def _is_dual_second(p: Procurement) -> bool:
    """是否为五选二二标段（含旧数据兼容：procurement_method=五选二 且 contract_section=二标段）."""
    if p.procurement_method == "五选二" and p.contract_section == "二标段":
        return True
    return False


def _is_dual_first(p: Procurement) -> bool:
    """是否为五选二一标段."""
    if p.procurement_method == "五选二" and p.contract_section == "一标段":
        return True
    return False


def _parse_form_data(form_data: str) -> dict:
    """Parse form_data JSON or repr to dict."""
    if not form_data:
        return {}
    try:
        return json.loads(form_data)
    except (json.JSONDecodeError, TypeError):
        try:
            import ast
            return ast.literal_eval(form_data) if form_data else {}
        except Exception:
            return {}


def _extract_remark(form_data: str | None) -> str:
    """从 form_data 读取备注（采购清单列）。"""
    d = _parse_form_data(form_data or "")
    if not isinstance(d, dict):
        return ""
    return str(d.get("remark") or "").strip()


def _merge_remark_into_form_data(form_data: str | None, remark: str) -> str:
    """将备注写入 form_data，保留其余字段。"""
    d = _parse_form_data(form_data or "")
    if not isinstance(d, dict):
        d = {}
    d = {**d, "remark": (remark or "").strip()}
    return json.dumps(d, ensure_ascii=False)


def _form_data_str_from_step2_time_records(step2: dict, time_records: list) -> str:
    """form_data：_time_records 不含「合同交底」行，其文本写入 hetong_jiaodi（目录.docx 仍用完整 time_records）。"""
    tr_out = []
    hj = str(step2.get("hetong_jiaodi") or "") if isinstance(step2, dict) else ""
    for r in time_records or []:
        fn = r.get("flow_name") or ""
        if fn == "合同交底":
            dv = r.get("date_val") or ""
            if dv:
                hj = dv
            continue
        tr_out.append({"flow_name": fn, "date_val": r.get("date_val", "")})
    merged = {**step2, "_time_records": tr_out, "hetong_jiaodi": hj}
    # 与前端 JSON.stringify 一致，避免首次保存时与 str(dict) repr 细微差异误判为「整表变更」而重生成全套流程文件
    return json.dumps(merged, ensure_ascii=False)


def _form_field_values_semantically_equal(a, b) -> bool:
    """比较表单标量：空值等价、数值宽松相等。"""
    def _is_blank(x) -> bool:
        return x is None or x == "" or (isinstance(x, str) and not str(x).strip())

    if _is_blank(a) and _is_blank(b):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        try:
            fa, fb = float(a), float(b)
            return fa == fb or math.isclose(fa, fb, rel_tol=0, abs_tol=1e-9)
        except (TypeError, ValueError):
            return False
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b)
    return a == b


def _non_time_form_fields_equal(old_fd: dict, new_fd: dict) -> bool:
    """除流程时间表相关字段外是否一致（用于判断是否仅更新目录.docx）。备注、空值形态、repr/JSON 差异不触发全套重生成。"""
    _skip = frozenset({"_time_records", "hetong_jiaodi", "remark"})
    keys = (set(old_fd.keys()) | set(new_fd.keys())) - _skip
    for k in keys:
        if not _form_field_values_semantically_equal(old_fd.get(k), new_fd.get(k)):
            return False
    return True


def _normalize_sign_date_str(val) -> str | None:
    """将签订日期统一为 YYYY-MM-DD 再比较，避免 2026-03-15 与 2026.3.15 误判为变更。"""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    s = s.replace(".", "-").replace("/", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if not m:
        return s
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"


def _sign_date_values_equal(a, b) -> bool:
    if _form_field_values_semantically_equal(a, b):
        return True
    na, nb = _normalize_sign_date_str(a), _normalize_sign_date_str(b)
    return na is not None and na == nb


def _normalize_time_records_for_compare(tr) -> list:
    """流程时间表行规范化后比较，避免空格/多余键导致首次保存误判。"""
    if not isinstance(tr, list):
        return []
    out = []
    for x in tr:
        if not isinstance(x, dict):
            continue
        out.append({
            "flow_name": str(x.get("flow_name") or "").strip(),
            "date_val": str(x.get("date_val") or "").strip(),
        })
    return out


def _is_blank_for_gain_detection(x) -> bool:
    """判断「是否视为未填写」，用于检测补填（空→有内容）。"""
    if x is None:
        return True
    if isinstance(x, str):
        return not str(x).strip()
    if isinstance(x, (list, dict)):
        return len(x) == 0
    return False


def _form_data_gained_nonblank_content(old_fd: dict, new_fd: dict) -> bool:
    """
    是否存在从空到有内容的补填。若有则不可跳过流程文件同步（须重生成 Word）。
    与 _form_payload_semantically_equal 并列使用，避免边界下语义比较漏判。
    """
    old_tr = _normalize_time_records_for_compare(old_fd.get("_time_records", []))
    new_tr = _normalize_time_records_for_compare(new_fd.get("_time_records", []))
    old_by_flow = {r["flow_name"]: r["date_val"] for r in old_tr}
    new_by_flow = {r["flow_name"]: r["date_val"] for r in new_tr}
    for fn, nd in new_by_flow.items():
        od = old_by_flow.get(fn, "")
        if _is_blank_for_gain_detection(od) and not _is_blank_for_gain_detection(nd):
            return True
    for fn, nd in new_by_flow.items():
        if fn not in old_by_flow and not _is_blank_for_gain_detection(nd):
            return True

    oh, nh = old_fd.get("hetong_jiaodi"), new_fd.get("hetong_jiaodi")
    if _is_blank_for_gain_detection(oh) and not _is_blank_for_gain_detection(nh):
        return True

    keys = set(old_fd.keys()) | set(new_fd.keys())
    skip_root = frozenset({"_time_records", "hetong_jiaodi"})
    for k in keys:
        if k in skip_root:
            continue
        ov, nv = old_fd.get(k), new_fd.get(k)
        if isinstance(ov, dict) and isinstance(nv, dict):
            if _form_data_gained_nonblank_content(ov, nv):
                return True
            continue
        if isinstance(ov, list) and isinstance(nv, list):
            if not ov and nv:
                return True
            continue
        if _is_blank_for_gain_detection(ov) and not _is_blank_for_gain_detection(nv):
            return True
    return False


def _form_payload_semantically_equal(old_fd: dict, new_fd: dict) -> bool:
    """整份 form_data 是否实质相同（含 _time_records 规范化、嵌套 dict 递归）。"""
    keys = set(old_fd.keys()) | set(new_fd.keys())
    for k in keys:
        ov, nv = old_fd.get(k), new_fd.get(k)
        if k == "sign_date":
            if not _sign_date_values_equal(ov, nv):
                return False
            continue
        if k == "_time_records":
            if _normalize_time_records_for_compare(ov) != _normalize_time_records_for_compare(nv):
                return False
            continue
        if isinstance(ov, dict) and isinstance(nv, dict):
            if not _form_payload_semantically_equal(ov, nv):
                return False
            continue
        if isinstance(ov, list) and isinstance(nv, list):
            if ov != nv:
                return False
            continue
        if not _form_field_values_semantically_equal(ov, nv):
            return False
    return True


def _procurement_file_sync_should_skip(
    old_form_data: str,
    new_form_data: str | None,
    old_suppliers: list,
    new_suppliers: list | None,
) -> bool:
    """相对上次保存，表单与供应商均无实质变化时跳过流程文件同步（避免首次打开即保存也全套覆盖）。"""
    if not new_form_data:
        return False
    if old_form_data == new_form_data:
        if new_suppliers is None:
            return True
    old_fd = _parse_form_data(old_form_data)
    new_fd = _parse_form_data(new_form_data)
    if not isinstance(old_fd, dict) or not isinstance(new_fd, dict):
        return (old_form_data or "").strip() == (new_form_data or "").strip() and new_suppliers is None
    # 补填空字段后必须重生成流程文件，不可因语义比较漏判而跳过
    if _form_data_gained_nonblank_content(old_fd, new_fd):
        return False
    if not _form_payload_semantically_equal(old_fd, new_fd):
        return False
    if new_suppliers is not None:
        def _norm_orm(s):
            return (getattr(s, "supplier_name", None), getattr(s, "quoted_price", None))
        def _norm_dict(s):
            return (s.get("supplier_name"), s.get("quoted_price"))
        old_norm = [_norm_orm(s) for s in old_suppliers]
        new_norm = [_norm_dict(s) for s in new_suppliers]
        if old_norm != new_norm:
            return False
    return True


def _overview_suppliers_payload(db: Session, proc: Procurement, suppliers: list) -> list:
    """采购总览：五选二共 5 家报价，仅两名中标分别显示「一标段」「二标段」，其余标段留空、中标留否。其它方式按库表。"""
    if proc.procurement_method != "五选二":
        return [
            {
                "supplier_name": s.supplier_name,
                "contact_person": s.contact_person,
                "contact_phone": s.contact_phone,
                "business_scope": s.business_scope or "",
                "tax_rate": s.tax_rate or "",
                "quoted_price": s.quoted_price or 0,
                "is_winner": bool(s.is_winner),
                "contract_section": s.contract_section or "",
            }
            for s in suppliers
        ]
    sibling = _find_dual_sibling(db, proc)
    if (proc.contract_section or "") == "一标段":
        proc_first, proc_second = proc, sibling
    elif (proc.contract_section or "") == "二标段":
        proc_first, proc_second = sibling, proc
    else:
        proc_first, proc_second = proc, sibling
    if not proc_first:
        proc_first = proc
    sups_a = (
        db.query(Supplier)
        .filter(Supplier.procurement_id == proc_first.id)
        .order_by(Supplier.rank)
        .all()
    )
    win_a = next((x for x in sups_a if x.is_winner), None)
    win_a_name = win_a.supplier_name if win_a else None
    win_b_name = None
    if proc_second:
        win_b = (
            db.query(Supplier)
            .filter(
                Supplier.procurement_id == proc_second.id,
                Supplier.is_winner == True,  # noqa: E712
            )
            .first()
        )
        win_b_name = win_b.supplier_name if win_b else None
    out = []
    for s in sups_a:
        sec, iw = "", False
        if win_a_name and s.supplier_name == win_a_name:
            sec, iw = "一标段", True
        elif win_b_name and s.supplier_name == win_b_name:
            sec, iw = "二标段", True
        out.append(
            {
                "supplier_name": s.supplier_name,
                "contact_person": s.contact_person,
                "contact_phone": s.contact_phone,
                "business_scope": s.business_scope or "",
                "tax_rate": s.tax_rate or "",
                "quoted_price": s.quoted_price or 0,
                "is_winner": iw,
                "contract_section": sec,
            }
        )
    return out


def _parent_winning_supplier_for_supplement(db: Session, parent: Procurement) -> Supplier | None:
    """补充协议与主合同中标一致；五选二主合同为某标段时取该标段对应排序名次供应商。"""
    sups = db.query(Supplier).filter(Supplier.procurement_id == parent.id).all()
    if not sups:
        return None
    sorted_sups = sorted(
        sups,
        key=lambda s: float(s.quoted_price) if s.quoted_price is not None and s.quoted_price != "" else float("inf"),
    )
    if parent.procurement_method == "五选二" and (parent.contract_section or "") == "二标段" and len(sorted_sups) > 1:
        return sorted_sups[1]
    return sorted_sups[0]


def _time_records_for_mulu(form_dict: dict) -> list:
    """目录.docx 用完整时间线：_time_records + 合同交底（来自 hetong_jiaodi 或行内日期）。"""
    if not isinstance(form_dict, dict):
        return []
    tr = [dict(x) for x in (form_dict.get("_time_records", []) or [])]
    hj = str(form_dict.get("hetong_jiaodi") or "").strip()
    if not any((r.get("flow_name") or "") == "合同交底" for r in tr):
        tr.append({"flow_name": "合同交底", "date_val": hj})
    else:
        for r in tr:
            if (r.get("flow_name") or "") == "合同交底" and not (str(r.get("date_val") or "").strip()) and hj:
                r["date_val"] = hj
    return tr


def _is_time_records_only_change(
    old_form_data: str,
    new_form_data: str | None,
    old_suppliers: list,
    new_suppliers: list | None,
) -> bool:
    """判断是否仅第4步流程时间表变更（_time_records + 合同交底 hetong_jiaodi）。是则只更新目录.docx。"""
    if not new_form_data:
        return False
    old_fd = _parse_form_data(old_form_data)
    new_fd = _parse_form_data(new_form_data)
    if not isinstance(old_fd, dict) or not isinstance(new_fd, dict):
        return False
    if not _non_time_form_fields_equal(old_fd, new_fd):
        return False
    old_tr = _normalize_time_records_for_compare(old_fd.get("_time_records", []))
    new_tr = _normalize_time_records_for_compare(new_fd.get("_time_records", []))
    old_hj = str(old_fd.get("hetong_jiaodi") or "").strip()
    new_hj = str(new_fd.get("hetong_jiaodi") or "").strip()
    if old_tr == new_tr and old_hj == new_hj:
        return False
    # 供应商未变更（new_suppliers 在 dump 中时需比较；不在 dump 则视为未变更）
    if new_suppliers is not None:
        def _norm_orm(s):
            return (getattr(s, "supplier_name", None), getattr(s, "quoted_price", None))
        def _norm_dict(s):
            return (s.get("supplier_name"), s.get("quoted_price"))
        old_norm = [_norm_orm(s) for s in old_suppliers]
        new_norm = [_norm_dict(s) for s in new_suppliers]
        if old_norm != new_norm:
            return False
    return True


def _find_dual_sibling(db: Session, proc: Procurement) -> Procurement | None:
    """Find the sibling procurement for 五选二 pair."""
    return resolve_dual_sibling(db, proc)


def _parse_supplement_seq(contract_number: str | None) -> int | None:
    """Parse supplement sequence from contract_number like 'xxx-补1' -> 1."""
    if not contract_number or "-补" not in contract_number:
        return None
    try:
        return int(contract_number.split("-补")[-1])
    except (ValueError, IndexError):
        return None


def _validate_quoted_price_vs_control(
    suppliers_data: list,
    control_price: float | None,
    method: str = "",
) -> None:
    """Validate all suppliers' quoted_price <= control price. 五选二：所有供应商都≤控制价."""
    if control_price is not None and control_price > 0:
        for i, s in enumerate(suppliers_data):
            qp = s.get("quoted_price") or 0
            if qp > control_price:
                name = s.get("supplier_name") or f"第{i + 1}家"
                raise HTTPException(
                    status_code=400,
                    detail=f"供应商「{name}」含税报价（{qp} 元）大于控制价（{control_price} 元），请调整。",
                )


def _validate_supplement_sign_dates(
    db: Session,
    parent_id: int,
    new_sign_date: str,
    *,
    exclude_proc_id: int | None = None,
    new_seq: int | None = None,
) -> None:
    """Validate supplement sign dates are in order: 补1 <= 补2 <= 补3..."""
    supplements = db.query(Procurement).filter(
        Procurement.parent_contract_id == parent_id,
        Procurement.procurement_method == "补充协议",
        Procurement.contract_number != "草稿",
    ).all()
    items = []
    for p in supplements:
        if p.id == exclude_proc_id:
            seq = _parse_supplement_seq(p.contract_number or "") or 0
            items.append((seq, new_sign_date))
        else:
            seq = _parse_supplement_seq(p.contract_number or "")
            if seq is not None:
                items.append((seq, p.sign_date or ""))
    if new_seq is not None:
        items.append((new_seq, new_sign_date))
    items.sort(key=lambda x: x[0])
    for i in range(1, len(items)):
        prev_date, curr_date = items[i - 1][1], items[i][1]
        if prev_date and curr_date and prev_date > curr_date:
            raise HTTPException(
                status_code=400,
                detail=f"同一主合同下，补充协议签订时间须按顺序递增：第{i}份（{prev_date}）应早于第{i + 1}份（{curr_date}），请调整。",
            )


@router.get("/parent-contract-options")
def get_parent_contract_options(
    project_id: int = Query(...),
    procurement_type: str = Query(...),
    is_dual_contract: bool = Query(False, description="主合同为五选二项目"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """补充协议主合同候选：同采购类型、非补充协议、合同编号已存在于台账。五选二时仅返回五选二产生的合同且已入台账。"""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    # 台账中存在的合同编号
    rows = db.query(Ledger.contract_number).filter(
        Ledger.project_id == project_id,
        Ledger.contract_number.isnot(None),
        Ledger.contract_number != "",
    ).distinct().all()
    ledger_cns = {r[0] for r in rows if r[0]}
    if not ledger_cns:
        return []
    q = db.query(Procurement).filter(
        Procurement.project_id == project_id,
        Procurement.procurement_type == procurement_type,
        Procurement.contract_number.isnot(None),
        Procurement.contract_number != "",
        Procurement.contract_number != "草稿",
        Procurement.procurement_method != "补充协议",
        Procurement.contract_number.in_(ledger_cns),
    )
    if is_dual_contract:
        q = q.filter(Procurement.procurement_method == "五选二")
    else:
        q = q.filter(Procurement.procurement_method != "五选二")
    items = q.order_by(Procurement.create_time.desc()).all()
    result = []
    for p in items:
        supp_count = db.query(Procurement).filter(
            Procurement.parent_contract_id == p.id,
        ).count()
        result.append({
            "id": p.id,
            "contract_number": p.contract_number,
            "project_name": p.project_name or "",
            "content": p.content or "",
            "contract_section": p.contract_section or "",
            "supplement_count": supp_count,
        })
    return result


@router.get("")
def list_procurements(
    project_id: int = Query(...),
    for_list: bool = Query(True, description="True=采购清单只显示1条五选二; False=归档用显示合并组"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """List procurements for a project."""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    items = db.query(Procurement).filter(Procurement.project_id == project_id).order_by(Procurement.create_time.desc()).all()
    result = []
    for p in items:
        sibling = None
        if _is_dual_second(p):
            if for_list:
                continue
            continue
        if _is_dual_first(p):
            sibling = _find_dual_sibling(db, p)

        suppliers = db.query(Supplier).filter(Supplier.procurement_id == p.id).order_by(Supplier.rank).all()
        # 按报价升序取第一名作为中选人（供应商/合同价随报价排名变更）
        sorted_by_price = sorted(
            suppliers,
            key=lambda s: (float(s.quoted_price) if s.quoted_price is not None and s.quoted_price != "" else float("inf")),
        )
        winner = sorted_by_price[0] if sorted_by_price else None
        remark = _extract_remark(p.form_data)

        # 补充协议合同价=新增金额；控制价用 form_data.supplement_control_price（非新增金额）
        if p.procurement_method == "补充协议":
            _contract_price = p.control_price or 0
            _control_price = None
            if p.form_data:
                try:
                    import json
                    try:
                        fd = json.loads(p.form_data)
                    except (json.JSONDecodeError, TypeError):
                        import ast
                        fd = ast.literal_eval(p.form_data) if p.form_data else {}
                    if isinstance(fd, dict):
                        sc = fd.get("supplement_control_price")
                        if sc is None or sc == "" or str(sc).strip() == "/":
                            _control_price = None
                        else:
                            try:
                                _control_price = float(sc)
                            except (ValueError, TypeError):
                                _control_price = None
                except Exception:
                    pass
        else:
            _contract_price = winner.quoted_price if winner else 0
            _control_price = p.control_price

        _supplier = winner.supplier_name if winner else ""
        _contract_price_display = None  # 五选二时用分行显示

        # 五选二：供应商、合同价分行显示；一标段=排序第1名+成交金额1，二标段=排序第2名+成交金额2
        if p.procurement_method == "五选二" and sibling:
            sorted_a = sorted(suppliers, key=lambda s: (float(s.quoted_price) if s.quoted_price is not None and s.quoted_price != "" else float("inf")))
            supp_b = db.query(Supplier).filter(Supplier.procurement_id == sibling.id).all()
            sorted_b = sorted(supp_b, key=lambda s: (float(s.quoted_price) if s.quoted_price is not None and s.quoted_price != "" else float("inf")))
            winner_a = sorted_a[0] if len(sorted_a) > 0 else None
            winner_b = sorted_b[1] if len(sorted_b) > 1 else (sorted_b[0] if sorted_b else None)
            supp_a_name = winner_a.supplier_name if winner_a else ""
            supp_b_name = winner_b.supplier_name if winner_b else ""
            _supplier = f"{supp_a_name}\n{supp_b_name}".strip() or ""
            # 从 form_data 取成交金额1、成交金额2
            cj1, cj2 = None, None
            if p.form_data:
                try:
                    import json
                    try:
                        fd = json.loads(p.form_data)
                    except (json.JSONDecodeError, TypeError):
                        import ast
                        fd = ast.literal_eval(p.form_data) if p.form_data else {}
                    if isinstance(fd, dict):
                        cj1 = fd.get("chengjiao_jine1")
                        cj2 = fd.get("chengjiao_jine2")
                except Exception:
                    pass
            if cj1 is None and winner_a:
                cj1 = winner_a.quoted_price
            if cj2 is None and winner_b:
                cj2 = winner_b.quoted_price
            c1 = (cj1 if cj1 is not None else 0) or 0
            c2 = (cj2 if cj2 is not None else 0) or 0
            _contract_price_display = f"{c1:,.2f}\n{c2:,.2f}"
            _contract_price = c1  # 保留首值供排序等

        item = {
            "id": p.id,
            "project_id": p.project_id,
            "procurement_type": p.procurement_type,
            "procurement_method": p.procurement_method,
            "project_name": p.project_name,
            "content": p.content,
            "control_price": _control_price,
            "contract_number": p.contract_number,
            "supplier": _supplier,
            "contract_price": _contract_price,
            "sign_date": p.sign_date,
            "remark": remark,
        }
        if _contract_price_display is not None:
            item["contract_price_display"] = _contract_price_display
        if not for_list and p.procurement_method == "五选二" and sibling:
            item["is_dual"] = True
            item["dual_ids"] = [p.id, sibling.id]
            cn_a = p.contract_number or ""
            cn_b = sibling.contract_number or ""
            item["contract_number"] = f"{cn_a}、{cn_b}" if cn_a and cn_b else (cn_a or cn_b)
        result.append(item)
    return result


@router.patch("/{procurement_id}/remark")
def patch_procurement_remark(
    procurement_id: int,
    body: ProcurementRemarkPatch,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """更新采购清单备注（写入 form_data.remark）。"""
    proc = db.query(Procurement).filter(Procurement.id == procurement_id).first()
    if not proc:
        raise HTTPException(status_code=404, detail="采购项目不存在")
    project = db.query(Project).filter(Project.id == proc.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    check_permission(project, current_user)
    new_fd = _merge_remark_into_form_data(proc.form_data, body.remark)
    proc.form_data = new_fd
    if _is_dual_first(proc):
        sibling = _find_dual_sibling(db, proc)
        if sibling:
            sibling.form_data = new_fd
    elif _is_dual_second(proc):
        sibling = _find_dual_sibling(db, proc)
        if sibling:
            sibling.form_data = new_fd
    db.commit()
    return {"message": "ok", "remark": (body.remark or "").strip()}


@router.get("/{procurement_id}")
def get_procurement(
    procurement_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Get procurement detail for overview (read-only)."""
    proc = db.query(Procurement).filter(Procurement.id == procurement_id).first()
    if not proc:
        raise HTTPException(status_code=404, detail="采购项目不存在")
    project = db.query(Project).filter(Project.id == proc.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    suppliers = db.query(Supplier).filter(Supplier.procurement_id == procurement_id).order_by(Supplier.rank).all()
    if proc.procurement_method == "五选二":
        sibling = _find_dual_sibling(db, proc)
        proc_first = proc if (proc.contract_section or "") == "一标段" else sibling
        if proc_first:
            suppliers = (
                db.query(Supplier)
                .filter(Supplier.procurement_id == proc_first.id)
                .order_by(Supplier.rank)
                .all()
            )
    step2 = {}
    time_records = []
    if proc.form_data:
        try:
            import ast
            import json
            try:
                data = json.loads(proc.form_data)
            except (json.JSONDecodeError, TypeError):
                data = ast.literal_eval(proc.form_data)
            if isinstance(data, dict):
                step2 = {k: v for k, v in data.items() if k != "_time_records"}
                time_records = list(data.get("_time_records", []) or [])
                hj = str(data.get("hetong_jiaodi") or "").strip()
                if not any((r.get("flow_name") or "") == "合同交底" for r in time_records):
                    time_records.append({"flow_name": "合同交底", "date_val": hj})
                elif hj:
                    for r in time_records:
                        if (r.get("flow_name") or "") == "合同交底" and not (r.get("date_val") or "").strip():
                            r["date_val"] = hj
        except Exception:
            pass
    # 工程项目信息从 Project 取（暂存载入时正确回填）；采购项目名称/内容若为草稿占位则返回空
    _proj_name = project.project_name or ""
    _proj_number = project.project_number or ""
    _proj_id = project.project_id or ""
    _construction = project.construction_unit or ""
    _total_price = project.total_contract_price or 0
    _proj_address = project.project_address or ""
    _department = project.department or ""
    _site_manager = project.site_manager or ""
    _site_manager_phone = project.site_manager_phone or ""
    _proc_name = proc.project_name if (proc.project_name or "") != "草稿" else ""
    _content = proc.content if (proc.content or "") != "草稿" else ""
    # 补充协议：总览「控制价」= supplement_control_price；新增金额单独字段 supplement_amount（与 DB proc.control_price 一致）
    _overview_control = proc.control_price
    _supplement_amount = None
    if proc.procurement_method == "补充协议":
        _supplement_amount = proc.control_price
        if isinstance(step2, dict):
            sc = step2.get("supplement_control_price")
            if sc is None or sc == "" or str(sc).strip() == "/":
                _overview_control = None
            else:
                try:
                    _overview_control = float(sc)
                except (ValueError, TypeError):
                    _overview_control = None
    return {
        "id": proc.id,
        "project_id": proc.project_id,
        "procurement_type": proc.procurement_type,
        "procurement_method": proc.procurement_method,
        "project_name": _proj_name,
        "project_number": _proj_number,
        "project_id_display": _proj_id,
        "construction_unit": _construction,
        "construction_contact_person": getattr(project, "construction_contact_person", None) or "",
        "construction_contact_phone": getattr(project, "construction_contact_phone", None) or "",
        "total_contract_price": _total_price,
        "project_address": _proj_address,
        "department": _department,
        "site_manager": _site_manager,
        "site_manager_phone": _site_manager_phone,
        "procurement_project_name": _proc_name,
        "content": _content,
        "control_price": _overview_control,
        "supplement_amount": _supplement_amount,
        "contract_number": proc.contract_number,
        "sign_date": proc.sign_date,
        "contract_section": proc.contract_section,
        "parent_contract_id": proc.parent_contract_id,
        "step2": step2,
        "suppliers": _overview_suppliers_payload(db, proc, suppliers),
        "time_records": time_records,
    }


@router.post("")
def create_procurement(
    data: ProcurementCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Create procurement with full flow."""
    project = db.query(Project).filter(Project.id == data.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    check_permission(project, current_user)

    step1 = data.step1.model_dump()
    step2 = data.step2.model_dump()
    step2["project_name"] = project.project_name
    step2["project_number"] = project.project_number
    step2["project_id"] = project.project_id
    step2["construction_unit"] = project.construction_unit
    step2["construction_contact_person"] = getattr(project, "construction_contact_person", None) or ""
    step2["construction_contact_phone"] = getattr(project, "construction_contact_phone", None) or ""
    step2["total_contract_price"] = project.total_contract_price
    step2["project_address"] = project.project_address
    step2["department"] = project.department
    step2["site_manager"] = project.site_manager
    step2["site_manager_phone"] = project.site_manager_phone
    # 补充协议控制价：前端可能在根级 supplement_control_price 传值，而 Pydantic 曾丢弃未声明字段；合并进 step2 以写入 form_data
    if step1["procurement_method"] == "补充协议":
        root_scp = getattr(data, "supplement_control_price", None)
        if root_scp is not None:
            step2["supplement_control_price"] = root_scp

    suppliers_data = [s.model_dump() for s in data.suppliers]
    method = step1["procurement_method"]
    is_draft = getattr(data, "is_draft", False) or False

    # 暂存：仅支持常规采购，放松校验
    if is_draft:
        if method == "五选二":
            raise HTTPException(status_code=400, detail="暂存不支持五选二")
        if method == "补充协议" and not data.parent_contract_id:
            raise HTTPException(status_code=400, detail="补充协议暂存需先选择主合同")
        if not suppliers_data:
            suppliers_data = [{"supplier_name": "", "contact_person": "", "contact_phone": "", "business_scope": "", "tax_rate": "", "quoted_price": 0}]
        if method == "邀请询比" and len(suppliers_data) < 3:
            while len(suppliers_data) < 3:
                suppliers_data.append({"supplier_name": "", "contact_person": "", "contact_phone": "", "business_scope": "", "tax_rate": "", "quoted_price": 0})
        if method == "补充协议" and len(suppliers_data) < 1:
            suppliers_data = [{"supplier_name": "", "contact_person": "", "contact_phone": "", "business_scope": "", "tax_rate": "", "quoted_price": 0}]
    else:
        # Validate supplier count
        if method in ["邀请询比"] and len(suppliers_data) < 3:
            raise HTTPException(status_code=400, detail="邀请询比需至少3家供应商")
        if method == "五选二" and len(suppliers_data) < 5:
            raise HTTPException(status_code=400, detail="五选二需至少5家供应商")
        if method in ["单一来源", "直接采购", "补充协议"] and len(suppliers_data) != 1:
            raise HTTPException(status_code=400, detail="该采购方式需1家供应商")

    # Sort suppliers by price
    sorted_suppliers = sorted(suppliers_data, key=lambda s: s.get("quoted_price", 0))

    if not is_draft and method != "补充协议":
        ctrl = step2.get("control_price")
        _validate_quoted_price_vs_control(
            suppliers_data,
            float(ctrl) if ctrl is not None and ctrl != "" else None,
            method,
        )

    if method == "五选二" and not is_draft:
        # Create two procurements
        form_data_str = _form_data_str_from_step2_time_records(step2, data.time_records)
        proc_a = Procurement(
            project_id=project.id,
            procurement_type=step1["procurement_type"],
            procurement_method=method,
            project_name=data.section_a_name or step2["procurement_project_name"],
            content=step2["content"],
            control_price=step2["control_price"],
            budget=round(step2["control_price"] / 10000),
            form_data=form_data_str,
            is_dual_contract=True,
            contract_section="一标段",
            contract_number="",
            sign_date=step2.get("sign_date", ""),
        )
        proc_b = Procurement(
            project_id=project.id,
            procurement_type=step1["procurement_type"],
            procurement_method=method,
            project_name=data.section_b_name or step2["procurement_project_name"],
            content=step2["content"],
            control_price=step2["control_price"],
            budget=round(step2["control_price"] / 10000),
            form_data=form_data_str,
            is_dual_contract=True,
            contract_section="二标段",
            contract_number="",
            sign_date=step2.get("sign_date", ""),
        )
        db.add(proc_a)
        db.add(proc_b)
        db.flush()

        # 五选二：一标段、二标段各存全部5家供应商，便于编辑时完整回显；一标段中选 rank1，二标段中选 rank2
        s1, s2 = sorted_suppliers[0], sorted_suppliers[1]
        for proc, winner_s, section in [(proc_a, s1, "一标段"), (proc_b, s2, "二标段")]:
            for i, s in enumerate(sorted_suppliers, 1):
                is_win = s == winner_s
                db.add(Supplier(
                    procurement_id=proc.id,
                    supplier_name=s["supplier_name"],
                    contact_person=s["contact_person"],
                    contact_phone=s["contact_phone"],
                    business_scope=s["business_scope"],
                    tax_rate=s["tax_rate"],
                    quoted_price=s["quoted_price"],
                    rank=i,
                    is_winner=is_win,
                    contract_section=section,
                ))
            # 台账在归档时生成，不在此创建

        # 五选二：创建采购项目专属文件夹（不分标段）并生成 Word 流程文件
        try:
            folder = create_procurement_folder(project, proc_a, step2.get("content") or step2.get("procurement_project_name") or "")
            add_procurement_ids_to_folder_marker(folder, [proc_a.id, proc_b.id])
            from app.services.word_service import _project_to_dict, build_context, get_template_path, render_docx
            ctx = build_context(
                _project_to_dict(project),
                {"funding_type": project.funding_type, "project_type": project.project_type, **step1},
                step2, suppliers_data, 0,
            )
            tpl_dir = get_template_path(project.funding_type, project.project_type, step1["procurement_type"], method)
            filenames = [d.name for d in tpl_dir.glob("*.docx")]
            for docx in tpl_dir.glob("*.docx"):
                render_docx(docx, ctx, folder / docx.name)
            title = f"{project.project_id}{project.project_name or ''}{step2.get('content', '')}"
            _save_mulu_docx(folder, title, data.time_records)
            reset_sync_status_records_for_new_procurement(db, proc_a.id)
            ensure_sync_status_records(db, proc_a.id, filenames + ["目录.docx"], folder)
        except Exception as e:
            logger.exception("五选二流程文件生成失败: %s", e)
            raise HTTPException(status_code=500, detail=f"流程文件生成失败: {str(e)}")

        db.commit()
        try:
            snap_sync_status_mtimes_from_disk(db, proc_a.id, folder)
        except Exception as e:
            logger.warning("五选二流程文件 mtime 提交后快照失败: %s", e)
        for pid in [proc_a.id, proc_b.id]:
            emit(EventType.PROCUREMENT_CREATED, {"project_id": project.id, "procurement_id": pid})
        return {"message": "ok", "procurement_ids": [proc_a.id, proc_b.id]}

    elif method == "补充协议" and is_draft:
        parent = db.query(Procurement).filter(Procurement.id == data.parent_contract_id).first()
        if not parent:
            raise HTTPException(status_code=400, detail="主合同不存在")
        supplement_content = data.supplement_content or step2.get("content", "")
        form_data_str = _form_data_str_from_step2_time_records(step2, data.time_records)
        proc = Procurement(
            project_id=project.id,
            procurement_type=step1["procurement_type"],
            procurement_method=method,
            parent_contract_id=data.parent_contract_id,
            project_name=step2.get("procurement_project_name") or parent.project_name or "",
            content=supplement_content or "",
            control_price=data.supplement_amount or 0,
            budget=0,
            form_data=form_data_str,
            contract_number="草稿",
            sign_date=step2.get("sign_date", ""),
            is_draft=True,
        )
        db.add(proc)
        db.flush()
        s = sorted_suppliers[0] if sorted_suppliers else {"supplier_name": "", "contact_person": "", "contact_phone": "", "business_scope": "", "tax_rate": "", "quoted_price": 0}
        parent_winner = _parent_winning_supplier_for_supplement(db, parent)
        sup_amt = data.supplement_amount or 0
        db.add(Supplier(
            procurement_id=proc.id,
            supplier_name=parent_winner.supplier_name if parent_winner else s.get("supplier_name", ""),
            contact_person=(parent_winner.contact_person if parent_winner else None) or s.get("contact_person", ""),
            contact_phone=(parent_winner.contact_phone if parent_winner else None) or s.get("contact_phone", ""),
            business_scope=(parent_winner.business_scope if parent_winner else None) or s.get("business_scope", ""),
            tax_rate=(parent_winner.tax_rate if parent_winner else None) or s.get("tax_rate", ""),
            quoted_price=sup_amt,
            rank=1,
            is_winner=True,
        ))
        # 补充协议暂存：创建专属文件夹（工程编号-草稿X）并生成全套流程文件（PRD 8.14/8.15）
        try:
            from app.services.word_service import _project_to_dict, build_context, get_template_path, render_docx
            folder_path = create_procurement_folder(project, proc, "草稿", db=db)
            supp_raw = getattr(data, "supplement_control_price", None)
            if supp_raw is None and isinstance(step2, dict):
                supp_raw = step2.get("supplement_control_price")
            supp_control = 0.0
            if supp_raw is not None and supp_raw != "" and str(supp_raw).strip() != "/":
                try:
                    supp_control = float(supp_raw)
                except (ValueError, TypeError):
                    supp_control = 0.0
            step2_supp = {**step2, "content": supplement_content, "control_price": supp_control, "project_name": parent.project_name or step2.get("procurement_project_name", "")}
            supp_suppliers = [{"supplier_name": parent_winner.supplier_name if parent_winner else s.get("supplier_name", ""), "contact_person": s.get("contact_person", ""), "contact_phone": s.get("contact_phone", ""), "business_scope": s.get("business_scope", ""), "tax_rate": s.get("tax_rate", ""), "quoted_price": sup_amt}]
            ctx = build_context(_project_to_dict(project), {"funding_type": project.funding_type, "project_type": project.project_type, **step1}, step2_supp, supp_suppliers, 0)
            tpl_dir = get_template_path(project.funding_type, project.project_type, step1["procurement_type"], method)
            filenames = [d.name for d in tpl_dir.glob("*.docx")]
            for docx in tpl_dir.glob("*.docx"):
                render_docx(docx, ctx, folder_path / docx.name)
            parent_contracts = db.query(Procurement).filter(
                Procurement.project_id == project.id,
                Procurement.parent_contract_id == data.parent_contract_id,
            ).all()
            supp_seq = len(parent_contracts)
            title = f"{project.project_id}{project.project_name or ''}{supplement_content or '补充协议'}补充协议{supp_seq}"
            _save_mulu_docx(folder_path, title, data.time_records)
            reset_sync_status_records_for_new_procurement(db, proc.id)
            ensure_sync_status_records(db, proc.id, filenames + ["目录.docx"], folder_path)
        except Exception as e:
            logger.exception("补充协议暂存流程文件生成失败: %s", e)
            raise HTTPException(status_code=500, detail=f"流程文件生成失败: {str(e)}")
        db.commit()
        try:
            snap_sync_status_mtimes_from_disk(db, proc.id, folder_path)
        except Exception as e:
            logger.warning("补充协议暂存 mtime 提交后快照失败: %s", e)
        db.refresh(proc)
        emit(EventType.PROCUREMENT_CREATED, {"project_id": project.id, "procurement_id": proc.id})
        return {"message": "ok", "procurement_id": proc.id}

    elif method == "补充协议" and not is_draft:
        parent = db.query(Procurement).filter(Procurement.id == data.parent_contract_id).first()
        if not parent:
            raise HTTPException(status_code=400, detail="主合同不存在")
        parent_winner = _parent_winning_supplier_for_supplement(db, parent)
        supplement_amount = data.supplement_amount or 0
        supplement_content = data.supplement_content or step2.get("content", "")
        supp_ctrl = getattr(data, "supplement_control_price", None)
        if supp_ctrl is None and isinstance(step2, dict):
            supp_ctrl = step2.get("supplement_control_price")
        if supp_ctrl is not None and supp_ctrl != "" and str(supp_ctrl).strip() != "/":
            try:
                ctrl_val = float(supp_ctrl)
                if ctrl_val > 0 and supplement_amount > ctrl_val:
                    raise HTTPException(
                        status_code=400,
                        detail=f"补充协议新增金额（{supplement_amount} 元）大于控制价（{ctrl_val} 元），请调整。",
                    )
            except (ValueError, TypeError):
                pass
        # 补充协议编号在合同文件归档时分配（主合同号-补X）
        parent_contracts = db.query(Procurement).filter(
            Procurement.project_id == project.id,
            Procurement.parent_contract_id == data.parent_contract_id,
        ).all()
        supp_seq = len(parent_contracts) + 1
        sign_date_val = step2.get("sign_date", "") or ""
        _validate_supplement_sign_dates(db, data.parent_contract_id, sign_date_val, new_seq=supp_seq)
        form_data_str = _form_data_str_from_step2_time_records(step2, data.time_records)
        proc = Procurement(
            project_id=project.id,
            procurement_type=step1["procurement_type"],
            procurement_method=method,
            parent_contract_id=data.parent_contract_id,
            project_name=step2.get("procurement_project_name") or parent.project_name or "",
            content=supplement_content,
            control_price=supplement_amount,
            budget=round(supplement_amount / 10000) if supplement_amount else 0,
            form_data=form_data_str,
            contract_number="",
            sign_date=sign_date_val,
        )
        db.add(proc)
        db.flush()
        s = sorted_suppliers[0]
        db.add(Supplier(
            procurement_id=proc.id,
            supplier_name=parent_winner.supplier_name if parent_winner else s["supplier_name"],
            contact_person=(parent_winner.contact_person if parent_winner else None) or s["contact_person"],
            contact_phone=(parent_winner.contact_phone if parent_winner else None) or s["contact_phone"],
            business_scope=(parent_winner.business_scope if parent_winner else None) or s.get("business_scope", ""),
            tax_rate=(parent_winner.tax_rate if parent_winner else None) or s.get("tax_rate", ""),
            quoted_price=supplement_amount,
            rank=1,
            is_winner=True,
        ))
        # 补充协议专属文件夹：常规 工程编号-采购内容-补X；五选二 工程编号-采购内容-标段名-补X（123.md）
        try:
            from app.config import PROCUREMENT_PROCESS_ROOT
            from app.services.word_service import _project_to_dict, build_context, get_template_path, render_docx
            content_safe = (supplement_content or parent.content or "").strip() or "补充协议"
            if parent.procurement_method == "五选二" and parent.contract_section:
                folder_name = f"{project.project_id}-{content_safe}-{parent.contract_section}-补{supp_seq}"
            else:
                folder_name = f"{project.project_id}-{content_safe}-补{supp_seq}"
            folder_path = PROCUREMENT_PROCESS_ROOT / f"{project.project_id} {project.project_name}" / folder_name
            folder_path.mkdir(parents=True, exist_ok=True)
            supp_control = 0.0
            if supp_ctrl is not None and supp_ctrl != "" and str(supp_ctrl).strip() != "/":
                try:
                    supp_control = float(supp_ctrl)
                except (ValueError, TypeError):
                    supp_control = 0.0
            step2_supp = {**step2, "content": supplement_content, "control_price": supp_control, "project_name": parent.project_name or step2.get("procurement_project_name", "")}
            supp_suppliers = [{"supplier_name": parent_winner.supplier_name if parent_winner else s["supplier_name"], "contact_person": s["contact_person"], "contact_phone": s["contact_phone"], "business_scope": s.get("business_scope", ""), "tax_rate": s.get("tax_rate", ""), "quoted_price": supplement_amount}]
            ctx = build_context(_project_to_dict(project), {"funding_type": project.funding_type, "project_type": project.project_type, **step1}, step2_supp, supp_suppliers, 0)
            tpl_dir = get_template_path(project.funding_type, project.project_type, step1["procurement_type"], method)
            filenames = [d.name for d in tpl_dir.glob("*.docx")]
            for docx in tpl_dir.glob("*.docx"):
                render_docx(docx, ctx, folder_path / docx.name)
            title = f"{project.project_id}{project.project_name or ''}{supplement_content or '补充协议'}补充协议{supp_seq}"
            _save_mulu_docx(folder_path, title, data.time_records)
            reset_sync_status_records_for_new_procurement(db, proc.id)
            ensure_sync_status_records(db, proc.id, filenames + ["目录.docx"], folder_path)
        except Exception as e:
            logger.exception("补充协议流程文件生成失败: %s", e)
            raise HTTPException(status_code=500, detail=f"流程文件生成失败: {str(e)}")
        db.commit()
        try:
            snap_sync_status_mtimes_from_disk(db, proc.id, folder_path)
        except Exception as e:
            logger.warning("补充协议 mtime 提交后快照失败: %s", e)
        emit(EventType.PROCUREMENT_CREATED, {"project_id": project.id, "procurement_id": proc.id})
        return {"message": "ok", "procurement_id": proc.id}

    else:
        # Single procurement (含暂存)
        contract_num = "草稿" if is_draft else ""
        form_data_str = _form_data_str_from_step2_time_records(step2, data.time_records)
        proc = Procurement(
            project_id=project.id,
            procurement_type=step1["procurement_type"],
            procurement_method=method,
            project_name=step2.get("procurement_project_name", "") or "",
            content=step2.get("content", "") or "",
            control_price=step2.get("control_price", 0) or 0,
            budget=round((step2.get("control_price", 0) or 0) / 10000),
            form_data=form_data_str,
            contract_number=contract_num,
            sign_date=step2.get("sign_date", ""),
            is_draft=is_draft,
        )
        db.add(proc)
        db.flush()
        winner = sorted_suppliers[0]
        for i, s in enumerate(sorted_suppliers, 1):
            db.add(Supplier(
                procurement_id=proc.id,
                supplier_name=s["supplier_name"],
                contact_person=s["contact_person"],
                contact_phone=s["contact_phone"],
                business_scope=s["business_scope"],
                tax_rate=s["tax_rate"],
                quoted_price=s["quoted_price"],
                rank=i,
                is_winner=(i == 1),
            ))
        # 台账在归档时生成，不在此创建

        # Create folder and generate Word (仅采购项目专属文件夹，不创建归档文件夹)
        try:
            folder = create_procurement_folder(project, proc, step2.get("content") or step2.get("procurement_project_name") or "草稿", db=db)
            from app.services.word_service import _project_to_dict, build_context, get_template_path, render_docx
            ctx = build_context(
                _project_to_dict(project),
                {"funding_type": project.funding_type, "project_type": project.project_type, **step1},
                step2, suppliers_data, 0
            )
            tpl_dir = get_template_path(project.funding_type, project.project_type, step1["procurement_type"], method)
            filenames = [d.name for d in tpl_dir.glob("*.docx")]
            for docx in tpl_dir.glob("*.docx"):
                render_docx(docx, ctx, folder / docx.name)
            title = f"{project.project_id}{project.project_name or ''}{step2.get('content', '')}"
            _save_mulu_docx(folder, title, data.time_records)
            reset_sync_status_records_for_new_procurement(db, proc.id)
            ensure_sync_status_records(db, proc.id, filenames + ["目录.docx"], folder)
        except Exception as e:
            logger.exception("采购项目流程文件生成失败: %s", e)
            raise HTTPException(status_code=500, detail=f"流程文件生成失败: {str(e)}")

        db.commit()
        try:
            snap_sync_status_mtimes_from_disk(db, proc.id, folder)
        except Exception as e:
            logger.warning("采购项目 mtime 提交后快照失败: %s", e)
        db.refresh(proc)
        emit(EventType.PROCUREMENT_CREATED, {"project_id": project.id, "procurement_id": proc.id})
        return {"message": "ok", "procurement_id": proc.id}


@router.put("/{procurement_id}")
def update_procurement(
    procurement_id: int,
    data: ProcurementUpdate,
    draft_only: bool = Query(False, description="仅暂存更新，不完成草稿"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Update procurement. 草稿完成时生成合同编号。draft_only=True 时仅更新表单不完成。"""
    proc = db.query(Procurement).filter(Procurement.id == procurement_id).first()
    if not proc:
        raise HTTPException(status_code=404, detail="采购项目不存在")
    project = db.query(Project).filter(Project.id == proc.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    check_permission(project, current_user)
    old_content = proc.content or ""
    old_project_name = proc.project_name or ""
    old_contract_number = proc.contract_number or ""
    dump = data.model_dump(exclude_unset=True)
    if draft_only and proc.is_draft:
        for k, v in dump.items():
            if k == "form_data":
                proc.form_data = v
            elif k == "project_name":
                proc.project_name = v
            elif k == "content":
                proc.content = v
            elif k == "supplement_amount" and v is not None:
                proc.control_price = v
            elif k == "supplement_content" and v is not None:
                proc.content = v
        if "suppliers" in dump and dump["suppliers"] is not None:
            db.query(Supplier).filter(Supplier.procurement_id == proc.id).delete()
            supp_amount = dump.get("supplement_amount") or 0
            _draft_supp = dump["suppliers"]
            _draft_sorted = sorted(_draft_supp, key=lambda s: float(s.get("quoted_price") or 0) if s.get("quoted_price") not in (None, "") else float("inf"))
            for i, s in enumerate(_draft_sorted, 1):
                qp = supp_amount if proc.procurement_method == "补充协议" and proc.parent_contract_id and i == 1 else s.get("quoted_price", 0)
                db.add(Supplier(
                    procurement_id=proc.id,
                    supplier_name=s.get("supplier_name", ""),
                    contact_person=s.get("contact_person", ""),
                    contact_phone=s.get("contact_phone", ""),
                    business_scope=s.get("business_scope", ""),
                    tax_rate=s.get("tax_rate", ""),
                    quoted_price=qp,
                    rank=i,
                    is_winner=(i == 1),
                ))
        # 暂存编辑：草稿文件夹保持 工程编号-草稿X，直到完成才重命名（PRD 8.15）
        db.commit()
        db.refresh(proc)
        emit(EventType.PROCUREMENT_UPDATED, {"project_id": project.id, "procurement_id": proc.id})
        return proc
    # 流程时间表单独更新：仅 _time_records 变更时只更新 目录.docx
    old_form_data = proc.form_data
    old_suppliers = list(db.query(Supplier).filter(Supplier.procurement_id == proc.id).order_by(Supplier.rank).all())

    if proc.is_draft and dump.get("form_data") and not draft_only:
        if proc.contract_number == "草稿":
            # 草稿转正式时不分配合同编号；按业务规则在合同文件上传归档时分配
            proc.contract_number = ""
        proc.is_draft = False
    for k, v in dump.items():
        if k == "form_data":
            proc.form_data = v
            # 五选二：同步更新兄弟标段的 form_data，避免编辑后打开仍为旧数据
            if proc.procurement_method == "五选二" and v:
                sibling = _find_dual_sibling(db, proc)
                if sibling:
                    sibling.form_data = v
            # 从 form_data 同步 sign_date、control_price 到 proc；五选二时同步到兄弟标段
            if v:
                try:
                    import json
                    try:
                        fd = json.loads(v)
                    except (json.JSONDecodeError, TypeError):
                        import ast
                        fd = ast.literal_eval(v)
                    if isinstance(fd, dict):
                        sign_val = fd.get("sign_date")
                        if sign_val is not None:
                            proc.sign_date = sign_val if isinstance(sign_val, str) else str(sign_val or "")
                            if proc.procurement_method == "五选二":
                                sib = _find_dual_sibling(db, proc)
                                if sib:
                                    sib.sign_date = proc.sign_date
                        if proc.procurement_method == "补充协议":
                            pjn = fd.get("procurement_project_name")
                            if pjn is not None and str(pjn).strip() != "":
                                proc.project_name = str(pjn).strip()
                        if proc.procurement_method != "补充协议" and "control_price" in fd:
                            ctrl_val = fd.get("control_price")
                            proc.control_price = float(ctrl_val) if ctrl_val is not None and ctrl_val != "" else None
                            if proc.procurement_method == "五选二":
                                sib = _find_dual_sibling(db, proc)
                                if sib:
                                    sib.control_price = proc.control_price
                except Exception:
                    pass
        elif k == "control_price" and proc.procurement_method != "补充协议":
            proc.control_price = float(v) if v is not None and v != "" else None
            if proc.procurement_method == "五选二":
                sib = _find_dual_sibling(db, proc)
                if sib:
                    sib.control_price = proc.control_price
        elif k == "suppliers" and v is not None:
            # 供应商数量需满足最低要求：邀请询比≥3，五选二≥5，其他=1
            min_supp = 5 if proc.procurement_method == "五选二" else (3 if proc.procurement_method == "邀请询比" else 1)
            if len(v) < min_supp:
                raise HTTPException(status_code=400, detail=f"{proc.procurement_method}需至少{min_supp}家供应商")
            if proc.procurement_method != "补充协议":
                fd = {}
                if proc.form_data:
                    try:
                        import json
                        try:
                            fd = json.loads(proc.form_data)
                        except (json.JSONDecodeError, TypeError):
                            import ast
                            fd = ast.literal_eval(proc.form_data) if proc.form_data else {}
                    except Exception:
                        pass
                ctrl = fd.get("control_price") if isinstance(fd, dict) else None
                if dump.get("form_data"):
                    try:
                        import json
                        try:
                            fd2 = json.loads(dump["form_data"])
                        except (json.JSONDecodeError, TypeError):
                            import ast
                            fd2 = ast.literal_eval(dump["form_data"])
                        if isinstance(fd2, dict):
                            ctrl = fd2.get("control_price", ctrl)
                    except Exception:
                        pass
                ctrl_f = float(ctrl) if ctrl is not None and ctrl != "" else None
                _validate_quoted_price_vs_control(v, ctrl_f, proc.procurement_method)
            db.query(Supplier).filter(Supplier.procurement_id == proc.id).delete()
            supp_amount = dump.get("supplement_amount") or 0
            # 询比/单源/五选二：按报价升序排序，中标列随供应商变更同步
            def _qp_key(s):
                q = s.get("quoted_price")
                if q is None or q == "" or (isinstance(q, str) and str(q).strip() == ""):
                    return float("inf")
                try:
                    return float(q)
                except (ValueError, TypeError):
                    return float("inf")
            sorted_v = sorted(v, key=_qp_key)
            sect = "一标段" if proc.procurement_method == "五选二" and proc.contract_section == "一标段" else None
            for i, s in enumerate(sorted_v, 1):
                qp = supp_amount if proc.procurement_method == "补充协议" and proc.parent_contract_id and i == 1 else s.get("quoted_price", 0)
                is_win = (i == 1) if proc.procurement_method != "五选二" else (i == 1 and sect == "一标段") or (i == 2 and sect != "一标段")
                sup = Supplier(
                    procurement_id=proc.id,
                    supplier_name=s.get("supplier_name", ""),
                    contact_person=s.get("contact_person", ""),
                    contact_phone=s.get("contact_phone", ""),
                    business_scope=s.get("business_scope", ""),
                    tax_rate=s.get("tax_rate", ""),
                    quoted_price=qp,
                    rank=i,
                    is_winner=is_win,
                )
                if sect:
                    sup.contract_section = sect
                db.add(sup)
            # 五选二：同步更新兄弟标段供应商（5家相同，一标段中选第1家、二标段中选第2家）
            if proc.procurement_method == "五选二" and len(sorted_v) >= 2:
                sibling = _find_dual_sibling(db, proc)
                if sibling:
                    db.query(Supplier).filter(Supplier.procurement_id == sibling.id).delete()
                    win_idx = 2 if sibling.contract_section == "二标段" else 1
                    sib_sect = sibling.contract_section or "二标段"
                    for i, s in enumerate(sorted_v, 1):
                        db.add(Supplier(
                            procurement_id=sibling.id,
                            supplier_name=s.get("supplier_name", ""),
                            contact_person=s.get("contact_person", ""),
                            contact_phone=s.get("contact_phone", ""),
                            business_scope=s.get("business_scope", ""),
                            tax_rate=s.get("tax_rate", ""),
                            quoted_price=s.get("quoted_price", 0),
                            rank=i,
                            is_winner=(i == win_idx),
                            contract_section=sib_sect,
                        ))
        elif k not in ("supplement_amount", "supplement_content"):
            setattr(proc, k, v)
    if "supplement_amount" in dump and dump["supplement_amount"] is not None:
        proc.control_price = dump["supplement_amount"]
    if "supplement_content" in dump and dump["supplement_content"] is not None:
        proc.content = dump["supplement_content"]
    if proc.procurement_method == "补充协议" and proc.parent_contract_id:
        supp_amt = proc.control_price or 0
        fd = {}
        if proc.form_data:
            try:
                import json
                try:
                    fd = json.loads(proc.form_data)
                except (json.JSONDecodeError, TypeError):
                    import ast
                    fd = ast.literal_eval(proc.form_data) if proc.form_data else {}
            except Exception:
                pass
        supp_ctrl = fd.get("supplement_control_price") if isinstance(fd, dict) else None
        if dump.get("form_data"):
            try:
                import json
                try:
                    fd2 = json.loads(dump["form_data"])
                except (json.JSONDecodeError, TypeError):
                    import ast
                    fd2 = ast.literal_eval(dump["form_data"])
                if isinstance(fd2, dict) and "supplement_control_price" in fd2:
                    supp_ctrl = fd2.get("supplement_control_price", supp_ctrl)
            except Exception:
                pass
        if supp_ctrl is not None and supp_ctrl != "" and str(supp_ctrl).strip() != "/":
            try:
                ctrl_val = float(supp_ctrl)
                if ctrl_val > 0 and supp_amt > ctrl_val:
                    raise HTTPException(
                        status_code=400,
                        detail=f"补充协议新增金额（{supp_amt} 元）大于控制价（{ctrl_val} 元），请调整。",
                    )
            except (ValueError, TypeError):
                pass
        _validate_supplement_sign_dates(db, proc.parent_contract_id, proc.sign_date or "", exclude_proc_id=proc.id)
    try:
        rename_procurement_folders(project, proc, old_content, old_project_name, old_contract_number, db)
    except Exception as e:
        logger.exception("采购项目文件夹重命名失败: %s", e)
        raise HTTPException(status_code=500, detail=f"文件夹重命名失败: {str(e)}")
    # 必须 flush 使新供应商写入 DB，否则 sync 查询不到，导致台账供应商/合同价空白
    db.flush()
    sync_ledgers_on_procurement_update(db, proc)
    if proc.procurement_method == "五选二":
        sibling = _find_dual_sibling(db, proc)
        if sibling:
            sync_ledgers_on_procurement_update(db, sibling)
    # 表单保存时触发流程文件同步（非草稿、非暂存）
    if not draft_only and not proc.is_draft:
        try:
            fd = {}
            if proc.form_data:
                try:
                    import json
                    fd = json.loads(proc.form_data)
                except (json.JSONDecodeError, TypeError):
                    import ast
                    fd = ast.literal_eval(proc.form_data) if proc.form_data else {}
            step1 = {
                "funding_type": project.funding_type or "工程类",
                "project_type": project.project_type or "集团内项目",
                "procurement_type": proc.procurement_type,
                "procurement_method": proc.procurement_method,
            }
            step2 = fd if isinstance(fd, dict) else {}
            if isinstance(step2, dict):
                step2.setdefault("project_name", project.project_name or "")
                step2.setdefault("project_number", project.project_number or "")
                step2.setdefault("project_id", project.project_id or "")
                step2.setdefault("construction_unit", project.construction_unit or "")
                step2.setdefault("construction_contact_person", getattr(project, "construction_contact_person", None) or "")
                step2.setdefault("construction_contact_phone", getattr(project, "construction_contact_phone", None) or "")
                step2.setdefault("total_contract_price", project.total_contract_price or 0)
                step2.setdefault("project_address", project.project_address or "")
                step2.setdefault("department", project.department or "")
                step2.setdefault("site_manager", project.site_manager or "")
                step2.setdefault("site_manager_phone", project.site_manager_phone or "")
            time_records = _time_records_for_mulu(step2) if isinstance(step2, dict) else []
            suppliers_data = [
                {"supplier_name": s.supplier_name, "contact_person": s.contact_person, "contact_phone": s.contact_phone,
                 "business_scope": s.business_scope or "", "tax_rate": s.tax_rate or "", "quoted_price": s.quoted_price or 0}
                for s in db.query(Supplier).filter(Supplier.procurement_id == proc.id).order_by(Supplier.rank).all()
            ]
            new_suppliers = dump.get("suppliers") if "suppliers" in dump else None
            if _procurement_file_sync_should_skip(
                old_form_data, proc.form_data, old_suppliers, new_suppliers
            ):
                pass  # 无实质变更：不触碰已生成的流程 Word
            else:
                time_records_only = _is_time_records_only_change(
                    old_form_data, proc.form_data, old_suppliers, new_suppliers
                )
                sync_process_files_on_form_save(
                    db, project, proc, step1, step2, suppliers_data, time_records,
                    time_records_only=time_records_only,
                )
        except Exception as e:
            logger.exception("流程文件同步失败 procurement_id=%s: %s", proc.id, e)
    db.commit()
    db.refresh(proc)
    emit(EventType.PROCUREMENT_UPDATED, {"project_id": project.id, "procurement_id": proc.id})
    return proc


@router.delete("/{procurement_id}")
def delete_procurement(
    procurement_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Delete procurement and sync folders, files, ledgers per PRD 4.3. For 五选二, deletes both 标段."""
    proc = db.query(Procurement).filter(Procurement.id == procurement_id).first()
    if not proc:
        raise HTTPException(status_code=404, detail="采购项目不存在")
    project = db.query(Project).filter(Project.id == proc.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    check_permission(project, current_user)
    to_delete = [proc]
    is_dual = proc.is_dual_contract or proc.procurement_method == "五选二"
    if is_dual:
        sibling = _find_dual_sibling(db, proc)
        if sibling:
            to_delete.append(sibling)
    # 五选二：在删除 DB 记录前先删除归档专属文件夹（PRD 8.13）
    if len(to_delete) == 2 and is_dual:
        folder_name = get_archive_contract_folder(project, to_delete[0], to_delete[1])
        if folder_name:
            archive_base = ARCHIVED_FILE_ROOT / f"{project.project_id}材料（设备）合同"
            archive_path = archive_base / folder_name
            if archive_path.exists():
                try:
                    shutil.rmtree(archive_path)
                except Exception as e:
                    logger.warning("五选二归档文件夹删除失败: %s", e)
    warnings = []
    for p in to_delete:
        try:
            delete_procurement_resources(db, project, p)
        except Exception as e:
            logger.exception("删除采购项目资源失败 procurement_id=%s: %s", p.id, e)
            warnings.append(f"采购项目 {p.id} 关联文件夹/文件删除失败: {str(e)}")
        db.query(ProcessFileSyncStatus).filter(
            ProcessFileSyncStatus.procurement_id == p.id,
        ).delete(synchronize_session=False)
        db.delete(p)
    db.commit()
    for p in to_delete:
        emit(EventType.PROCUREMENT_DELETED, {"project_id": project.id, "procurement_id": p.id})
    res = {"message": "ok"}
    if warnings:
        res["warnings"] = warnings
    return res
