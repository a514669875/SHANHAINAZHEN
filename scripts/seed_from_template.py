#!/usr/bin/env python3
"""
按 docs/seed_data_template.md 灌入 3 个工程项目及各自 56 条采购（含五选二、多轮补充协议）。

用法（项目根目录）:
  python scripts/seed_from_template.py              # 若工程编号已存在则跳过或报错
  python scripts/seed_from_template.py --clear      # 先删除三个模板工程再重新灌数
  python scripts/seed_from_template.py --clear-only # 仅删除三个模板工程，不灌数、不写文件
  python scripts/seed_from_template.py --seed 42    # 指定随机种子，复现「50 条流程文件」抽样

可选参数：
  --process-files-count N   随机 N 条采购拷贝模板目录下全套 .docx（默认 50，0 关闭）
  --no-process-files        不拷贝流程文件（仅归档占位）
  --seed INT                抽样随机种子（默认每次不同）

灌数完成后（只要库中存在三个模板工程之一）会：
  - 为每条采购在正式归档目录写入 1 份「占位归档合同」PDF（已存在同名记录则跳过）
  - 对已具备归档占位合同的采购按业务规则生成/补全台账（与上传归档逻辑一致：五选二需两段均有占位）
  - 随机抽取若干条采购，将对应 Word 模板目录内全部 .docx 拷入流程文件夹并登记 sync 状态

需已执行 init_db.py，且存在 admin 用户。
"""
from __future__ import annotations

import argparse
import random
import re
import sys
from datetime import date
from pathlib import Path
from shutil import copy2

from sqlalchemy import not_

backend_dir = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(backend_dir))

from app.database import SessionLocal  # noqa: E402
from app.models.user import User  # noqa: E402
from app.models.project import Project  # noqa: E402
from app.models.procurement import Procurement  # noqa: E402
from app.models.supplier import Supplier  # noqa: E402
from app.models.ledger import Ledger  # noqa: E402
from app.services.project_service import create_project_folders  # noqa: E402
from app.models.file import File  # noqa: E402
from app.config import ARCHIVED_FILE_ROOT  # noqa: E402
from app.services.archive_service import get_archive_contract_folder  # noqa: E402
from app.services.process_file_sync_service import ensure_sync_status_records  # noqa: E402
from app.services.procurement_service import (  # noqa: E402
    create_ledger_record,
    create_procurement_folder,
    generate_contract_number,
    get_next_contract_seq,
    get_process_folder_for_procurement,
    sync_ledgers_on_procurement_update,
)
from app.services.word_service import get_template_path  # noqa: E402

# 与模板 B 节一致（工程编号用于 --clear）
TEMPLATE_PROJECT_IDS = ("专土23-07", "专机23-17N", "日常25-21")

PROJECT_DEFS = [
    {
        "funding_type": "工程类",
        "project_type": "集团内项目",
        "project_number": "1023000401",
        "project_id": "专土23-07",
        "project_name": "广钢新城加压站建设工程总承包",
        "department": "广钢项目部",
        "site_manager": "林洁盛",
        "site_manager_phone": "15875914450",
        "construction_unit": "广州市自来水有限公司",
        "construction_contact_person": "张三",
        "construction_contact_phone": "15965488745",
        "total_contract_price": 257834771.1,
        "project_duration": "暂定2023年12月28日-2025年12月31日",
        "funding_source": "部分自有资金、部分财政资金",
        "project_address": "广州市荔湾区",
    },
    {
        "funding_type": "工程类",
        "project_type": "集团外项目",
        "project_number": "1023001302Z",
        "project_id": "专机23-17N",
        "project_name": "惠州至肇庆高速公路白云至三水段项目给排水迁改工程",
        "department": "项目中心",
        "site_manager": "柯少玲",
        "site_manager_phone": "19925952895",
        "construction_unit": "广州惠肇高速有限公司",
        "construction_contact_person": "李四",
        "construction_contact_phone": "15999954874",
        "total_contract_price": 279577385.0,
        "project_duration": "24个月",
        "funding_source": "企业自筹",
        "project_address": "惠州至肇庆高速公路",
    },
    {
        "funding_type": "自有资金",
        "project_type": "",
        "project_number": "/",
        "project_id": "日常25-21",
        "project_name": "2025年度安全文明物资采购项目",
        "department": "工程综合部",
        "site_manager": "严桂焕",
        "site_manager_phone": "15999954874",
        "construction_unit": "广州自来水专业建安有限公司",
        "construction_contact_person": "王五",
        "construction_contact_phone": "15999965587",
        "total_contract_price": 2000000.0,
        "project_duration": "/",
        "funding_source": "自有资金",
        "project_address": "广州市荔湾区",
    },
]

HY3 = "广州建材有限公司 100000; 南方建材公司 105000; 东方建材厂 110000"
W5 = "日成公司 180000; 至高建设 180600; 山和集团 181740; 宏佳建设 182000; 振邦贸易 183000"


def _norm_date(s: str) -> str:
    s = (s or "").strip().replace(" ", "")
    parts = re.split(r"[./-]", s)
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        return f"{y:04d}-{m:02d}-{d:02d}"
    return s or ""


def parse_supplier_list(raw: str) -> list[tuple[str, float]]:
    """名称 + 金额；分号分隔；金额取段末数字（支持无空格粘连）。负数按原样入库。"""
    if not raw or not str(raw).strip():
        return [("", 0.0)]
    out: list[tuple[str, float]] = []
    for chunk in str(raw).split(";"):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.search(r"([-+]?[\d.]+)\s*$", chunk)
        if not m:
            out.append((chunk, 0.0))
            continue
        price_s = m.group(1)
        name = chunk[: m.start()].strip()
        try:
            price = float(price_s)
        except ValueError:
            price = 0.0
        out.append((name or "供应商", price))
    return out if out else [("", 0.0)]


def next_supplement_index(db, parent_id: int) -> int:
    parent = db.query(Procurement).filter(Procurement.id == parent_id).first()
    if not parent or not parent.contract_number:
        return 1
    base = parent.contract_number
    mx = 0
    for p in db.query(Procurement).filter(Procurement.parent_contract_id == parent_id).all():
        cn = p.contract_number or ""
        if not cn.startswith(base) or "-补" not in cn:
            continue
        mm = re.search(r"-补(\d+)", cn)
        if mm:
            mx = max(mx, int(mm.group(1)))
    return mx + 1


def add_suppliers_for_proc(
    db,
    procurement_id: int,
    raw: str,
    *,
    dual_section: str | None = None,
    winner_rank: int | None = None,
):
    pairs = parse_supplier_list(raw)
    for i, (name, price) in enumerate(pairs, 1):
        win = False
        if dual_section:
            if dual_section == "一标段" and winner_rank is not None:
                win = i == winner_rank
            elif dual_section == "二标段" and winner_rank is not None:
                win = i == winner_rank
        else:
            win = i == 1
        db.add(
            Supplier(
                procurement_id=procurement_id,
                supplier_name=name,
                quoted_price=price,
                rank=i,
                is_winner=win,
                contract_section=dual_section or None,
            )
        )


def build_procurement_specs_56() -> list[dict]:
    """
    56 条采购规格，与模板 C 节一致。
    parent: None | 'm1'..'m4b' | 'd1'..'d4b' | 'e1'..'e4b' | 'l1'..'l4b'
    """
    rows: list[dict] = []

    def R(
        ptype: str,
        method: str,
        pname: str,
        content: str,
        ctrl: float,
        sig: str,
        sup: str,
        *,
        dual: bool = False,
        parent: str | None = None,
    ):
        rows.append(
            {
                "ptype": ptype,
                "method": method,
                "pname": pname,
                "content": content,
                "control": float(ctrl),
                "sign": _norm_date(sig),
                "sup": sup,
                "dual": dual,
                "parent": parent,
            }
        )

    # --- 材料采购 1–14 ---
    R("材料采购", "邀请询比", "混凝土", "混凝土", 413182.78, "2025.6.16", HY3)
    R("材料采购", "单一来源", "钢筋", "钢筋", 2348572.78, "2025.9.20", "独家供应商 480000")
    R("材料采购", "直接采购", "水泥", "水泥", 37966.5, "2025.10.15", "某某贸易公司 37966.5")
    R("材料采购", "五选二", "阀门", "阀门", 11379249.14, "2024.4.17", W5, dual=True)
    R("材料采购", "补充协议", "混凝土", "混凝土", 126589, "2025.7.13", "广州建材有限公司 20000", parent="m1")
    R("材料采购", "补充协议", "钢筋", "钢筋", 41312.78, "2025.10.30", "独家供应商 80000", parent="m2")
    R("材料采购", "补充协议", "水泥", "水泥", 238572, "2025.11.26", "某某贸易公司 13444.15", parent="m3")
    R("材料采购", "补充协议", "阀门", "阀门", 0, "2024.6.23", "日成公司 80555", parent="m4a")
    R("材料采购", "补充协议", "阀门", "阀门", 113749.14, "2024.6.23", "至高建设 60066", parent="m4b")
    R("材料采购", "补充协议", "混凝土", "混凝土", 325642.12, "2025.8.5", "广州建材有限公司 30000", parent="m1")
    R("材料采购", "补充协议", "钢筋", "钢筋", 13182.78, "2025.11.24", "独家供应商 13066.22", parent="m2")
    R("材料采购", "补充协议", "水泥", "水泥", 238572.78, "2025.12.12", "某某贸易公司 22556.77", parent="m3")
    R("材料采购", "补充协议", "阀门", "阀门", 0, "2024.7.4", "日成公司 -523", parent="m4a")
    R("材料采购", "补充协议", "阀门", "阀门", 392149.14, "2024.7.4", "至高建设 55633.2", parent="m4b")

    # --- 设备采购 15–28 ---
    R("设备采购", "邀请询比", "水泵", "水泵", 413182.78, "2025.6.16", HY3)
    R("设备采购", "单一来源", "水表", "水表", 2348572.78, "2025.9.20", "独家供应商 480000")
    R("设备采购", "直接采购", "电设备", "电设备", 37966.5, "2025.10.15", "某某贸易公司 37966.5")
    R("设备采购", "五选二", "供水设备", "供水设备", 11379249.14, "2024.4.17", W5, dual=True)
    R("设备采购", "补充协议", "水泵", "水泵", 126589, "2025.7.13", "广州建材有限公司 20000", parent="d1")
    R("设备采购", "补充协议", "水表", "水表", 41312.78, "2025.10.30", "独家供应商 80000", parent="d2")
    R("设备采购", "补充协议", "电设备", "电设备", 238572, "2025.11.26", "某某贸易公司 13444.15", parent="d3")
    R("设备采购", "补充协议", "供水设备", "供水设备", 0, "2024.6.23", "日成公司 80555", parent="d4a")
    R("设备采购", "补充协议", "供水设备", "供水设备", 113749.14, "2024.6.23", "至高建设 60066", parent="d4b")
    R("设备采购", "补充协议", "水泵", "水泵", 325642.12, "2025.8.5", "广州建材有限公司 30000", parent="d1")
    R("设备采购", "补充协议", "水表", "水表", 13182.78, "2025.11.24", "独家供应商 13066.22", parent="d2")
    R("设备采购", "补充协议", "电设备", "电设备", 238572.78, "2025.12.12", "某某贸易公司 22556.77", parent="d3")
    R("设备采购", "补充协议", "供水设备", "供水设备", 0, "2024.7.4", "日成公司 -523", parent="d4a")
    R("设备采购", "补充协议", "供水设备", "供水设备", 392149.14, "2024.7.4", "至高建设 55633.2", parent="d4b")

    # --- 机械租赁 29–42 ---
    R("机械租赁", "邀请询比", "洒水车", "洒水车", 413182.78, "2025.6.16", HY3)
    R("机械租赁", "单一来源", "挖机", "挖机", 2348572.78, "2025.9.20", "独家供应商 480000")
    R("机械租赁", "直接采购", "吊机", "吊机", 37966.5, "2025.10.15", "某某贸易公司 37966.5")
    R("机械租赁", "五选二", "起重机", "起重机", 11379249.14, "2024.4.17", W5, dual=True)
    R("机械租赁", "补充协议", "洒水车", "洒水车", 126589, "2025.7.13", "广州建材有限公司 20000", parent="e1")
    R("机械租赁", "补充协议", "挖机", "挖机", 41312.78, "2025.10.30", "独家供应商 80000", parent="e2")
    R("机械租赁", "补充协议", "吊机", "吊机", 238572, "2025.11.26", "某某贸易公司 13444.15", parent="e3")
    R("机械租赁", "补充协议", "起重机", "起重机", 0, "2024.6.23", "日成公司 80555", parent="e4a")
    R("机械租赁", "补充协议", "起重机", "起重机", 113749.14, "2024.6.23", "至高建设 60066", parent="e4b")
    R("机械租赁", "补充协议", "洒水车", "洒水车", 325642.12, "2025.8.5", "广州建材有限公司 30000", parent="e1")
    R("机械租赁", "补充协议", "挖机", "挖机", 13182.78, "2025.11.24", "独家供应商 13066.22", parent="e2")
    R("机械租赁", "补充协议", "吊机", "吊机", 238572.78, "2025.12.12", "某某贸易公司 22556.77", parent="e3")
    R("机械租赁", "补充协议", "起重机", "起重机", 0, "2024.7.4", "日成公司 -523", parent="e4a")
    R("机械租赁", "补充协议", "起重机", "起重机", 392149.14, "2024.7.4", "至高建设 55633.2", parent="e4b")

    # --- 材料租赁 43–56 ---
    R("材料租赁", "邀请询比", "钢板桩", "钢板桩", 413182.78, "2025.6.16", HY3)
    R("材料租赁", "单一来源", "机器具", "机器具", 2348572.78, "2025.9.20", "独家供应商 480000")
    R("材料租赁", "直接采购", "顶管架", "顶管架", 37966.5, "2025.10.15", "某某贸易公司 37966.5")
    R("材料租赁", "五选二", "模板", "模板", 11379249.14, "2024.4.17", W5, dual=True)
    R("材料租赁", "补充协议", "钢板桩", "钢板桩", 126589, "2025.7.13", "广州建材有限公司 20000", parent="l1")
    R("材料租赁", "补充协议", "机器具", "机器具", 41312.78, "2025.10.30", "独家供应商 80000", parent="l2")
    R("材料租赁", "补充协议", "顶管架", "顶管架", 238572, "2025.11.26", "某某贸易公司 13444.15", parent="l3")
    R("材料租赁", "补充协议", "模板", "模板", 0, "2024.6.23", "日成公司 80555", parent="l4a")
    R("材料租赁", "补充协议", "模板", "模板", 113749.14, "2024.6.23", "至高建设 60066", parent="l4b")
    R("材料租赁", "补充协议", "钢板桩", "钢板桩", 325642.12, "2025.8.5", "广州建材有限公司 30000", parent="l1")
    R("材料租赁", "补充协议", "机器具", "机器具", 13182.78, "2025.11.24", "独家供应商 13066.22", parent="l2")
    R("材料租赁", "补充协议", "顶管架", "顶管架", 238572.78, "2025.12.12", "某某贸易公司 22556.77", parent="l3")
    R("材料租赁", "补充协议", "模板", "模板", 0, "2024.7.4", "日成公司 -523", parent="l4a")
    R("材料租赁", "补充协议", "模板", "模板", 392149.14, "2024.7.4", "至高建设 55633.2", parent="l4b")

    assert len(rows) == 56, len(rows)
    return rows


def clear_template_projects(db) -> int:
    n = 0
    for pid in TEMPLATE_PROJECT_IDS:
        p = db.query(Project).filter(Project.project_id == pid).first()
        if p:
            db.delete(p)
            n += 1
    db.commit()
    return n


def create_project_row(db, officer_ids: str, d: dict) -> Project:
    p = Project(
        funding_type=d["funding_type"],
        project_type=d.get("project_type") or "",
        project_number=d["project_number"],
        project_id=d["project_id"],
        project_name=d["project_name"],
        department=d["department"],
        site_manager=d["site_manager"],
        site_manager_phone=d["site_manager_phone"],
        construction_unit=d["construction_unit"],
        construction_contact_person=d.get("construction_contact_person") or "",
        construction_contact_phone=d.get("construction_contact_phone") or "",
        total_contract_price=d["total_contract_price"],
        project_duration=d.get("project_duration") or "",
        funding_source=d["funding_source"],
        project_address=d.get("project_address") or "",
        procurement_officers=officer_ids,
        create_date=date.today(),
    )
    db.add(p)
    db.flush()
    try:
        create_project_folders(p)
    except ValueError as e:
        print(f"  工程 {p.project_id} 文件夹: {e}")
    db.commit()
    db.refresh(p)
    return p


def seed_one_project(db, project: Project, specs: list[dict]) -> tuple[int, dict[str, Procurement]]:
    refs: dict[str, Procurement] = {}
    created = 0

    def reg_m4_dual(pa: Procurement, pb: Procurement):
        refs["m4a"] = pa
        refs["m4b"] = pb

    def reg_d4_dual(pa: Procurement, pb: Procurement):
        refs["d4a"] = pa
        refs["d4b"] = pb

    def reg_e4_dual(pa: Procurement, pb: Procurement):
        refs["e4a"] = pa
        refs["e4b"] = pb

    def reg_l4_dual(pa: Procurement, pb: Procurement):
        refs["l4a"] = pa
        refs["l4b"] = pb

    idx = 0
    while idx < len(specs):
        sp = specs[idx]
        ptype = sp["ptype"]
        method = sp["method"]

        if sp["dual"]:
            seq = get_next_contract_seq(db, project.id, ptype)
            pa = Procurement(
                project_id=project.id,
                procurement_type=ptype,
                procurement_method=method,
                project_name=f"{sp['pname']}（一标段）",
                content=sp["content"],
                control_price=sp["control"],
                is_dual_contract=True,
                contract_section="一标段",
                contract_number=generate_contract_number(project, ptype, seq),
                sign_date=sp["sign"],
                is_draft=False,
                form_data="{}",
            )
            pb = Procurement(
                project_id=project.id,
                procurement_type=ptype,
                procurement_method=method,
                project_name=f"{sp['pname']}（二标段）",
                content=sp["content"],
                control_price=sp["control"],
                is_dual_contract=True,
                contract_section="二标段",
                contract_number=generate_contract_number(project, ptype, seq + 1),
                sign_date=sp["sign"],
                is_draft=False,
                form_data="{}",
            )
            db.add(pa)
            db.add(pb)
            db.flush()
            add_suppliers_for_proc(db, pa.id, sp["sup"], dual_section="一标段", winner_rank=1)
            add_suppliers_for_proc(db, pb.id, sp["sup"], dual_section="二标段", winner_rank=2)
            try:
                create_procurement_folder(project, pa, sp["content"])
            except Exception:
                pass
            db.commit()
            db.refresh(pa)
            db.refresh(pb)
            # 登记 ref（按段）
            if ptype == "材料采购":
                reg_m4_dual(pa, pb)
            elif ptype == "设备采购":
                reg_d4_dual(pa, pb)
            elif ptype == "机械租赁":
                reg_e4_dual(pa, pb)
            elif ptype == "材料租赁":
                reg_l4_dual(pa, pb)
            created += 2
            idx += 1
            continue

        if sp["parent"]:
            key = sp["parent"]
            parent = refs.get(key)
            if not parent:
                raise RuntimeError(f"缺少父采购 ref {key}，工程 {project.project_id}")
            sup_idx = next_supplement_index(db, parent.id)
            cn = f"{parent.contract_number}-补{sup_idx}"
            proc = Procurement(
                project_id=project.id,
                procurement_type=ptype,
                procurement_method=method,
                parent_contract_id=parent.id,
                project_name=sp["pname"],
                content=sp["content"],
                control_price=sp["control"],
                contract_number=cn,
                sign_date=sp["sign"],
                is_draft=False,
                form_data="{}",
            )
            db.add(proc)
            db.flush()
            pairs = parse_supplier_list(sp["sup"])
            wname, wprice = pairs[0] if pairs else ("", 0.0)
            db.add(
                Supplier(
                    procurement_id=proc.id,
                    supplier_name=wname,
                    quoted_price=wprice,
                    rank=1,
                    is_winner=True,
                )
            )
            try:
                from app.config import PROCUREMENT_PROCESS_ROOT

                folder = PROCUREMENT_PROCESS_ROOT / f"{project.project_id} {project.project_name}" / f"{project.project_id}-{sp['content']}-补{sup_idx}"
                folder.mkdir(parents=True, exist_ok=True)
            except Exception:
                pass
            db.commit()
            db.refresh(proc)
            created += 1
            idx += 1
            continue

        # 主合同（非五选二、非补充）
        seq = get_next_contract_seq(db, project.id, ptype)
        proc = Procurement(
            project_id=project.id,
            procurement_type=ptype,
            procurement_method=method,
            project_name=sp["pname"],
            content=sp["content"],
            control_price=sp["control"],
            contract_number=generate_contract_number(project, ptype, seq),
            sign_date=sp["sign"],
            is_draft=False,
            form_data="{}",
        )
        db.add(proc)
        db.flush()
        add_suppliers_for_proc(db, proc.id, sp["sup"])
        try:
            create_procurement_folder(project, proc, sp["content"])
        except Exception:
            pass
        db.commit()
        db.refresh(proc)

        # 登记主合同 ref
        if ptype == "材料采购" and method == "邀请询比" and sp["content"] == "混凝土":
            refs["m1"] = proc
        elif ptype == "材料采购" and method == "单一来源" and sp["content"] == "钢筋":
            refs["m2"] = proc
        elif ptype == "材料采购" and method == "直接采购" and sp["content"] == "水泥":
            refs["m3"] = proc
        elif ptype == "设备采购" and method == "邀请询比" and sp["content"] == "水泵":
            refs["d1"] = proc
        elif ptype == "设备采购" and method == "单一来源" and sp["content"] == "水表":
            refs["d2"] = proc
        elif ptype == "设备采购" and method == "直接采购" and sp["content"] == "电设备":
            refs["d3"] = proc
        elif ptype == "机械租赁" and method == "邀请询比" and sp["content"] == "洒水车":
            refs["e1"] = proc
        elif ptype == "机械租赁" and method == "单一来源" and sp["content"] == "挖机":
            refs["e2"] = proc
        elif ptype == "机械租赁" and method == "直接采购" and sp["content"] == "吊机":
            refs["e3"] = proc
        elif ptype == "材料租赁" and method == "邀请询比" and sp["content"] == "钢板桩":
            refs["l1"] = proc
        elif ptype == "材料租赁" and method == "单一来源" and sp["content"] == "机器具":
            refs["l2"] = proc
        elif ptype == "材料租赁" and method == "直接采购" and sp["content"] == "顶管架":
            refs["l3"] = proc

        created += 1
        idx += 1

    return created, refs


# 最小合法 PDF，供归档占位（体积极小）
MINIMAL_PDF_BYTES = b"""%PDF-1.4
1 0 obj<<>>endobj
trailer<<>>
%%EOF
"""


def find_dual_sibling_seed(db, proc: Procurement) -> Procurement | None:
    """五选二另一标段：同工程、同采购内容、同签约日（多种五选二并存时与 find_dual_sibling 任意 first 区分）。"""
    if proc.procurement_method != "五选二":
        return None
    want = "二标段" if proc.contract_section == "一标段" else "一标段"
    return (
        db.query(Procurement)
        .filter(
            Procurement.project_id == proc.project_id,
            Procurement.procurement_method == "五选二",
            Procurement.content == proc.content,
            Procurement.sign_date == proc.sign_date,
            Procurement.contract_section == want,
            Procurement.id != proc.id,
        )
        .first()
    )


def get_primary_procurement_id_seed(db, proc: Procurement) -> int:
    if proc.procurement_method != "五选二":
        return proc.id
    if proc.contract_section == "一标段":
        return proc.id
    sib = find_dual_sibling_seed(db, proc)
    return sib.id if sib else proc.id


def placeholder_archive_pdf_name(proc: Procurement) -> str:
    if proc.procurement_method == "五选二" and proc.contract_section:
        return f"占位归档合同_{proc.contract_section}.pdf"
    return "占位归档合同.pdf"


def resolve_archive_contract_folder_name(db, project: Project, proc: Procurement) -> str:
    if proc.procurement_method == "五选二":
        sib = find_dual_sibling_seed(db, proc)
        if not sib:
            return proc.contract_number or ""
        return get_archive_contract_folder(project, proc, sib)
    return proc.contract_number or ""


def create_placeholder_archive_contract(db, project: Project, proc: Procurement, admin_id: int) -> bool:
    """在正式归档目录写入占位 PDF 并插入 File 行。已存在同名文件记录则返回 False。"""
    if proc.is_draft or not proc.contract_number or proc.contract_number == "草稿":
        return False
    fname = placeholder_archive_pdf_name(proc)
    if db.query(File).filter(File.procurement_id == proc.id, File.file_name == fname).first():
        return False
    folder_name = resolve_archive_contract_folder_name(db, project, proc)
    if not folder_name:
        return False
    archive_base = f"{project.project_id}材料（设备）合同"
    rel = f"{archive_base}/{folder_name}"
    full_dir = ARCHIVED_FILE_ROOT / rel
    full_dir.mkdir(parents=True, exist_ok=True)
    full_path = full_dir / fname
    full_path.write_bytes(MINIMAL_PDF_BYTES)
    rel_path = f"{rel}/{fname}"
    db.add(
        File(
            procurement_id=proc.id,
            file_name=fname,
            file_path=rel_path,
            file_size=len(MINIMAL_PDF_BYTES),
            file_type="合同文件",
            is_contract=True,
            print_mode="单面",
            owner_user_id=admin_id,
            storage_computer_ip="",
            storage_computer_name="",
            storage_path=str(full_path.resolve()),
            upload_by=admin_id,
        )
    )
    return True


def _parse_seed_form_data(form_data: str) -> dict:
    if not form_data:
        return {}
    try:
        import json

        return json.loads(form_data)
    except Exception:
        try:
            import ast

            return ast.literal_eval(form_data) if form_data else {}
        except Exception:
            return {}


def _procurement_has_archived_contract_file(db, procurement_id: int) -> bool:
    """非 _temp 路径的合同文件（含种子直接写入归档的占位 PDF）。"""
    return (
        db.query(File)
        .filter(
            File.procurement_id == procurement_id,
            File.is_contract.is_(True),
            not_(File.file_path.like("_temp/%")),
        )
        .first()
        is not None
    )


def ensure_ledgers_for_archived_placeholder_contracts(db) -> int:
    """
    与 try_archive_and_create_ledgers 中台账规则对齐：已有归档路径合同占位时补建 Ledger。
    五选二需一、二标段均有归档占位合同后再为两段各建台账。
    返回本次新建的台账条数。
    """
    pairs = collect_template_project_procurements(db)
    n_new = 0
    dual_done: set[frozenset[int]] = set()

    for project, proc in pairs:
        if proc.procurement_method == "五选二":
            sib = find_dual_sibling_seed(db, proc)
            if not sib:
                continue
            pf, ps = (proc, sib) if proc.contract_section == "一标段" else (sib, proc)
            key = frozenset({pf.id, ps.id})
            if key in dual_done:
                continue
            dual_done.add(key)
            if not _procurement_has_archived_contract_file(db, pf.id) or not _procurement_has_archived_contract_file(
                db, ps.id
            ):
                continue
            fd = _parse_seed_form_data(pf.form_data or "")
            cj1, cj2 = fd.get("chengjiao_jine1"), fd.get("chengjiao_jine2")
            for p, sect, cp_val in [(pf, "一标段", cj1), (ps, "二标段", cj2)]:
                if db.query(Ledger).filter(Ledger.procurement_id == p.id).first():
                    continue
                supps = db.query(Supplier).filter(Supplier.procurement_id == p.id).all()
                sorted_supps = sorted(
                    supps,
                    key=lambda s: (
                        float(s.quoted_price) if s.quoted_price is not None and s.quoted_price != "" else float("inf")
                    ),
                )
                rank_idx = 0 if sect == "一标段" else 1
                if rank_idx >= len(sorted_supps):
                    continue
                winner = sorted_supps[rank_idx]
                cp = cp_val
                if cp is None or cp == "":
                    cp = (
                        float(winner.quoted_price)
                        if winner.quoted_price is not None and winner.quoted_price != ""
                        else 0.0
                    )
                else:
                    cp = float(cp)
                create_ledger_record(db, project, p, winner.supplier_name or "", cp, p.sign_date or "")
                n_new += 1
            db.flush()
            for p in (pf, ps):
                f_row = (
                    db.query(File)
                    .filter(
                        File.procurement_id == p.id,
                        File.is_contract.is_(True),
                        not_(File.file_path.like("_temp/%")),
                    )
                    .order_by(File.id)
                    .first()
                )
                if f_row:
                    lg = db.query(Ledger).filter(Ledger.procurement_id == p.id).first()
                    if lg:
                        lg.pdf_preview_path = f"/api/files/{f_row.id}/content"
            for p in (pf, ps):
                sync_ledgers_on_procurement_update(db, p)
            continue

        if not _procurement_has_archived_contract_file(db, proc.id):
            continue

        existing_ledger = db.query(Ledger).filter(Ledger.procurement_id == proc.id).first()
        f_row = (
            db.query(File)
            .filter(
                File.procurement_id == proc.id,
                File.is_contract.is_(True),
                not_(File.file_path.like("_temp/%")),
            )
            .order_by(File.id)
            .first()
        )
        if existing_ledger:
            if f_row and not (existing_ledger.pdf_preview_path or "").strip():
                existing_ledger.pdf_preview_path = f"/api/files/{f_row.id}/content"
            continue

        winner = db.query(Supplier).filter(Supplier.procurement_id == proc.id).order_by(Supplier.rank).first()
        if not winner:
            continue
        parent_num = None
        if proc.parent_contract_id:
            parent = db.query(Procurement).filter(Procurement.id == proc.parent_contract_id).first()
            parent_num = parent.contract_number if parent else None
        if proc.procurement_method == "补充协议":
            contract_price = float(proc.control_price or 0)
            fd = _parse_seed_form_data(proc.form_data or "")
            supp_ctrl = fd.get("supplement_control_price")
            supp_ctrl_val = float(supp_ctrl) if supp_ctrl is not None and supp_ctrl != "" else None
            create_ledger_record(
                db,
                project,
                proc,
                winner.supplier_name or "",
                contract_price,
                proc.sign_date or "",
                parent_num,
                control_price_override=supp_ctrl_val,
            )
        else:
            qp = winner.quoted_price
            contract_price = float(qp) if qp is not None and qp != "" else 0.0
            create_ledger_record(
                db, project, proc, winner.supplier_name or "", contract_price, proc.sign_date or "", parent_num
            )

        if proc.parent_contract_id and parent_num:
            parent_ledger = db.query(Ledger).filter(Ledger.procurement_id == proc.parent_contract_id).first()
            if parent_ledger:
                existing_supp = (parent_ledger.supplement_contracts or "").strip()
                parent_ledger.supplement_contracts = (
                    f"{existing_supp}\n{proc.contract_number}".strip() if existing_supp else proc.contract_number
                )
        n_new += 1
        db.flush()
        if f_row:
            lg = db.query(Ledger).filter(Ledger.procurement_id == proc.id).first()
            if lg:
                lg.pdf_preview_path = f"/api/files/{f_row.id}/content"
        sync_ledgers_on_procurement_update(db, proc)

    return n_new


def ensure_mulu_docx(folder: Path) -> None:
    """流程目录若无 目录.docx 则生成极简占位。"""
    mulu = folder / "目录.docx"
    if mulu.exists():
        return
    from docx import Document

    doc = Document()
    doc.add_paragraph("目录（种子占位）")
    doc.save(mulu)


def copy_full_process_docx_bundle(db, project: Project, proc: Procurement) -> None:
    """将模板目录下全部 .docx 拷入该采购流程文件夹，并补 目录.docx、sync 状态。"""
    tpl = get_template_path(
        project.funding_type,
        project.project_type,
        proc.procurement_type,
        proc.procurement_method,
    )
    folder = get_process_folder_for_procurement(project, proc, db)
    folder.mkdir(parents=True, exist_ok=True)
    if tpl.is_dir():
        for f in sorted(tpl.glob("*.docx")):
            if f.name.startswith("~$"):
                continue
            try:
                copy2(f, folder / f.name)
            except OSError as e:
                print(f"    警告: 拷贝失败 {f.name}: {e}")
    else:
        print(f"    警告: 模板目录不存在 {tpl}，仅写入目录占位")
    ensure_mulu_docx(folder)
    filenames = sorted(p.name for p in folder.glob("*.docx"))
    if not filenames:
        filenames = ["目录.docx"]
    primary_id = get_primary_procurement_id_seed(db, proc)
    ensure_sync_status_records(db, primary_id, filenames, folder)


def collect_template_project_procurements(db) -> list[tuple[Project, Procurement]]:
    projects = (
        db.query(Project)
        .filter(Project.project_id.in_(TEMPLATE_PROJECT_IDS))
        .order_by(Project.id)
        .all()
    )
    pairs: list[tuple[Project, Procurement]] = []
    for pr in projects:
        for proc in (
            db.query(Procurement)
            .filter(Procurement.project_id == pr.id)
            .order_by(Procurement.id)
            .all()
        ):
            pairs.append((pr, proc))
    return pairs


def seed_archive_placeholders_and_process_files(
    db,
    admin_id: int,
    *,
    process_files_count: int,
    rng: random.Random,
) -> tuple[int, int, int]:
    """返回 (新建归档占位条数, 实际写入流程文件的采购条数, 本次新建台账条数)。"""
    pairs = collect_template_project_procurements(db)
    if not pairs:
        return 0, 0, 0
    n_new_arch = 0
    for pr, proc in pairs:
        if create_placeholder_archive_contract(db, pr, proc, admin_id):
            n_new_arch += 1
    db.flush()
    n_ledgers = ensure_ledgers_for_archived_placeholder_contracts(db)
    db.commit()

    n_proc_written = 0
    if process_files_count and process_files_count > 0:
        k = min(process_files_count, len(pairs))
        picked = rng.sample(pairs, k)
        for pr, proc in picked:
            try:
                copy_full_process_docx_bundle(db, pr, proc)
                n_proc_written += 1
            except Exception as e:
                print(f"    警告: 流程文件写入失败 procurement_id={proc.id}: {e}")
        db.commit()
    return n_new_arch, n_proc_written, n_ledgers


def run_clear_only() -> None:
    """仅删除三个模板工程（不创建、不灌数、不写占位与流程）。"""
    db = SessionLocal()
    try:
        n = clear_template_projects(db)
        db.commit()
        print(f"已仅删除模板工程 {n} 条（专土23-07 / 专机23-17N / 日常25-21），未灌数、未写占位与台账。")
    finally:
        db.close()


def run_seed(
    clear_first: bool,
    *,
    process_files_count: int = 50,
    random_seed: int | None = None,
) -> None:
    specs = build_procurement_specs_56()
    db = SessionLocal()
    try:
        admin = db.query(User).filter(User.username == "admin").first()
        if not admin:
            raise RuntimeError("请先运行 init_db.py 创建管理员用户")
        officer_ids = str(admin.id)

        if clear_first:
            n = clear_template_projects(db)
            print(f"已删除模板工程记录 {n} 条（级联采购等）")

        total_proc = 0
        for pdef in PROJECT_DEFS:
            existing = db.query(Project).filter(Project.project_id == pdef["project_id"]).first()
            if existing and not clear_first:
                print(f"跳过已存在工程: {pdef['project_id']}")
                continue
            if existing and clear_first:
                # clear 已删，不应存在
                pass
            proj = create_project_row(db, officer_ids, pdef)
            n, _ = seed_one_project(db, proj, specs)
            total_proc += n
            print(f"  工程 {proj.project_id} 已创建，采购条数: {n}")

        print("=" * 50)
        if total_proc:
            print(
                f"完成。本次写入采购记录 {total_proc} 条（每工程 60 条；"
                "模板 56 条规格中含 4 个五选二各拆成 2 条数据库记录）"
            )
        else:
            print("完成。未写入新数据（工程已存在时请加 --clear 覆盖，或换工程编号）")

        tpl_pairs = collect_template_project_procurements(db)
        rng = random.Random(random_seed) if random_seed is not None else random.Random()
        n_arch, n_pf, n_led = seed_archive_placeholders_and_process_files(
            db,
            admin.id,
            process_files_count=process_files_count,
            rng=rng,
        )
        if tpl_pairs:
            print(
                f"归档合同占位：本次新建 {n_arch} 条 File 记录（其余采购若已有占位则跳过）"
            )
            print(f"台账：本次新建 {n_led} 条（已有台账的采购跳过；五选二需两段均有归档占位后生成）")
            if process_files_count and process_files_count > 0:
                k = min(process_files_count, len(tpl_pairs))
                print(
                    f"流程文件：自 {len(tpl_pairs)} 条采购中随机抽 {k} 条，已成功写入 {n_pf} 条（模板目录无文件时仅含目录占位）"
                )
            else:
                print("流程文件：已跳过（--process-files-count 0 或 --no-process-files）")
        print("=" * 50)
    finally:
        db.close()


def main():
    ap = argparse.ArgumentParser(description="按 seed_data_template 灌入测试数据")
    ap.add_argument("--clear", action="store_true", help="先删除专土23-07/专机23-17N/日常25-21 三个工程再灌数")
    ap.add_argument(
        "--clear-only",
        action="store_true",
        help="仅删除上述三个模板工程，不灌数、不写占位合同与流程文件（与 --clear 互斥用法：不要同时使用）",
    )
    ap.add_argument(
        "--process-files-count",
        type=int,
        default=50,
        metavar="N",
        help="随机 N 条采购拷贝流程模板下全部 docx（默认 50；0 关闭）",
    )
    ap.add_argument(
        "--no-process-files",
        action="store_true",
        help="不拷贝流程文件（等价于 --process-files-count 0）",
    )
    ap.add_argument(
        "--seed",
        type=int,
        default=None,
        metavar="INT",
        help="随机抽样种子，指定后每次抽样相同（默认每次不同）",
    )
    args = ap.parse_args()
    if args.clear_only:
        if args.clear:
            print("提示: 已指定 --clear-only，将只执行删除，忽略 --clear。")
        run_clear_only()
        return
    pfc = 0 if args.no_process_files else args.process_files_count
    run_seed(args.clear, process_files_count=pfc, random_seed=args.seed)


if __name__ == "__main__":
    main()
